# -*- coding: utf-8 -*-
"""独立的 XLSX/CSV 总表拆分与分表合并功能。"""

from __future__ import annotations

import csv
import io
import os
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell, WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


TABLE_ACTION_SPLIT = "split"
TABLE_ACTION_MERGE = "merge"
OUTPUT_DIR_NAME = "生成结果"
OUTPUT_SUFFIX = "_输出结果"


class TableOperationError(Exception):
    """可以直接显示给使用者的拆分/合并错误。"""


class TableOperationCancelled(TableOperationError):
    """使用者主动取消拆分或合并。"""


@dataclass(frozen=True)
class TableOperationResult:
    action: str
    generated_rows: int
    output_file_count: int
    source_file_count: int
    output_dir: str
    output_path: str | None
    output_paths: tuple[str, ...]


ProgressCallback = Callable[[str, int, int, str], None]
CancelCheck = Callable[[], bool]


def _check_cancel(cancel_check: CancelCheck | None) -> None:
    if cancel_check and cancel_check():
        raise TableOperationCancelled("处理已取消。")


def _report_progress(
    callback: ProgressCallback | None,
    phase: str,
    completed: int,
    total: int,
    message: str,
) -> None:
    if callback:
        callback(phase, max(0, completed), max(0, total), message)


def _remove_created_files(paths: Iterable[Path]) -> None:
    for path in paths:
        try:
            if path.exists():
                path.unlink()
        except OSError:
            pass


def _is_nonempty(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or value != "")


def _last_nonempty_row(rows: Sequence[Sequence[object]]) -> int:
    for index in range(len(rows) - 1, -1, -1):
        if any(_is_nonempty(value) for value in rows[index]):
            return index + 1
    return 0


def _last_nonempty_column(values: Sequence[object]) -> int:
    for index in range(len(values) - 1, -1, -1):
        if _is_nonempty(values[index]):
            return index + 1
    return 0


def _unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    counter = 1
    while True:
        candidate = path.with_name(f"{path.stem}({counter}){path.suffix}")
        if not candidate.exists():
            return candidate
        counter += 1


def _split_output_paths(source: Path, output_dir: Path, count: int) -> list[Path]:
    stems = [f"{index:03d}_{source.stem}{OUTPUT_SUFFIX}" for index in range(1, count + 1)]
    plain = [output_dir / f"{stem}{source.suffix.lower()}" for stem in stems]
    if not any(path.exists() for path in plain):
        return plain
    version = 1
    while True:
        candidates = [
            output_dir / f"{stem}({version}){source.suffix.lower()}" for stem in stems
        ]
        if not any(path.exists() for path in candidates):
            return candidates
        version += 1


def _decode_csv(path: Path) -> tuple[str, list[list[str]]]:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = content.decode(encoding)
            rows = list(csv.reader(io.StringIO(text, newline="")))
            return encoding, rows
        except UnicodeDecodeError:
            continue
    raise TableOperationError(f"无法识别 CSV 文件编码：{path.name}")


def _trim_csv_rows(rows: list[list[str]]) -> list[list[str]]:
    last_row = _last_nonempty_row(rows)
    return rows[:last_row]


def _validate_csv_table(path: Path, rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    rows = _trim_csv_rows(rows)
    if not rows:
        raise TableOperationError(f"文件为空，未找到表头：{path.name}")
    header_width = _last_nonempty_column(rows[0])
    if header_width == 0:
        raise TableOperationError(f"第一行表头为空：{path.name}")
    header = rows[0][:header_width]
    data: list[list[str]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if any(_is_nonempty(value) for value in row[header_width:]):
            raise TableOperationError(
                f"{path.name} 第 {row_number} 行在表头范围之外还有数据，请先补全表头。"
            )
        data.append((row + [""] * header_width)[:header_width])
    return header, data


def _write_csv(path: Path, rows: Iterable[Sequence[object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerows(rows)


def _first_visible_sheet(workbook, path: Path) -> Worksheet:
    for worksheet in workbook.worksheets:
        if worksheet.sheet_state == "visible":
            return worksheet
    raise TableOperationError(f"工作簿没有可见工作表：{path.name}")


def _worksheet_bounds(worksheet: Worksheet, path: Path) -> tuple[int, int]:
    last_row = 0
    last_column = 0
    for row in worksheet.iter_rows():
        row_has_value = False
        for cell in row:
            if _is_nonempty(cell.value):
                row_has_value = True
                last_column = max(last_column, cell.column)
        if row_has_value:
            last_row = row[0].row
    if last_row == 0:
        raise TableOperationError(f"工作表为空，未找到表头：{path.name}")
    header_width = 0
    for cell in worksheet[1]:
        if _is_nonempty(cell.value):
            header_width = cell.column
    if header_width == 0:
        raise TableOperationError(f"第一行表头为空：{path.name}")
    if last_column > header_width:
        for row in worksheet.iter_rows(
            min_row=2, max_row=last_row, min_col=header_width + 1, max_col=last_column
        ):
            for cell in row:
                if _is_nonempty(cell.value):
                    raise TableOperationError(
                        f"{path.name} 第 {cell.row} 行第 {cell.column} 列超出表头范围，"
                        "请先补全表头。"
                    )
    return last_row, header_width


def _header_values(worksheet: Worksheet, width: int) -> tuple[object, ...]:
    return tuple(worksheet.cell(1, column).value for column in range(1, width + 1))


def _copy_cell(source: Cell, target: Cell) -> None:
    target.value = source.value
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def _copy_row_dimension(source: Worksheet, target: Worksheet, source_row: int, target_row: int) -> None:
    source_dimension = source.row_dimensions[source_row]
    target_dimension = target.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    target_dimension.font = copy(source_dimension.font)
    target_dimension.fill = copy(source_dimension.fill)
    target_dimension.border = copy(source_dimension.border)
    target_dimension.alignment = copy(source_dimension.alignment)
    target_dimension.protection = copy(source_dimension.protection)
    target_dimension.number_format = source_dimension.number_format


def _copy_sheet_settings(source: Worksheet, target: Worksheet, width: int) -> None:
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_view.zoomScale = source.sheet_view.zoomScale
    target.freeze_panes = source.freeze_panes
    for column in range(1, width + 1):
        key = get_column_letter(column)
        source_dimension = source.column_dimensions[key]
        target_dimension = target.column_dimensions[key]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.bestFit = source_dimension.bestFit
        target_dimension.outlineLevel = source_dimension.outlineLevel
        target_dimension.collapsed = source_dimension.collapsed
        target_dimension.font = copy(source_dimension.font)
        target_dimension.fill = copy(source_dimension.fill)
        target_dimension.border = copy(source_dimension.border)
        target_dimension.alignment = copy(source_dimension.alignment)
        target_dimension.protection = copy(source_dimension.protection)
        target_dimension.number_format = source_dimension.number_format


def _create_xlsx_output(
    source: Worksheet,
    path: Path,
    source_rows: Sequence[int],
    width: int,
    completed_before: int = 0,
    total_data_rows: int = 0,
    progress_message: str = "正在生成分表",
    progress_callback: ProgressCallback | None = None,
    cancel_check: CancelCheck | None = None,
) -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = source.title
    _copy_sheet_settings(source, target, width)
    row_map = {source_row: target_row for target_row, source_row in enumerate(source_rows, start=1)}
    copied_data_rows = 0
    for source_row, target_row in row_map.items():
        _copy_row_dimension(source, target, source_row, target_row)
        for column in range(1, width + 1):
            _copy_cell(source.cell(source_row, column), target.cell(target_row, column))
        if source_row != 1:
            copied_data_rows += 1
            if copied_data_rows % 2_000 == 0:
                _check_cancel(cancel_check)
                _report_progress(
                    progress_callback,
                    "split",
                    completed_before + copied_data_rows,
                    total_data_rows,
                    progress_message,
                )

    for merged_range in source.merged_cells.ranges:
        rows = range(merged_range.min_row, merged_range.max_row + 1)
        if all(row in row_map for row in rows):
            target.merge_cells(
                start_row=row_map[merged_range.min_row],
                start_column=merged_range.min_col,
                end_row=row_map[merged_range.max_row],
                end_column=merged_range.max_col,
            )
        elif any(row in row_map for row in rows):
            workbook.close()
            raise TableOperationError(
                f"工作表存在跨越拆分边界的合并单元格：{merged_range}，请先取消合并。"
            )

    if source.auto_filter.ref:
        target.auto_filter.ref = f"A1:{get_column_letter(width)}{len(source_rows)}"
    _check_cancel(cancel_check)
    try:
        workbook.save(path)
    finally:
        workbook.close()


def split_table(
    source_path: Path,
    max_rows: int,
    status_callback: Callable[[str], None] | None = None,
    progress_callback: ProgressCallback | None = None,
    cancel_check: CancelCheck | None = None,
) -> TableOperationResult:
    source = source_path.resolve()
    if not source.is_file():
        raise TableOperationError("请选择存在的总表文件。")
    suffix = source.suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise TableOperationError("拆分总表只支持 .xlsx 和 .csv 文件。")
    if max_rows < 2:
        raise TableOperationError("每个分表最大行数必须至少为 2（包含表头）。")

    output_dir = source.parent / OUTPUT_DIR_NAME
    output_dir.mkdir(parents=True, exist_ok=True)
    data_per_file = max_rows - 1

    output_paths: list[Path] = []
    try:
        _check_cancel(cancel_check)
        _report_progress(progress_callback, "scan", 0, 0, "正在读取总表……")
        if suffix == ".csv":
            _encoding, raw_rows = _decode_csv(source)
            header, data_rows = _validate_csv_table(source, raw_rows)
            if not data_rows:
                raise TableOperationError("总表只有表头，没有可拆分的数据。")
            count = (len(data_rows) + data_per_file - 1) // data_per_file
            output_paths = _split_output_paths(source, output_dir, count)
            generated_rows = len(data_rows)
            _report_progress(
                progress_callback, "split", 0, generated_rows, "准备生成分表……"
            )
            for index, path in enumerate(output_paths):
                _check_cancel(cancel_check)
                if status_callback:
                    status_callback(f"正在生成第 {index + 1}/{count} 个分表……")
                start = index * data_per_file
                end = min(generated_rows, start + data_per_file)
                _write_csv(path, [header, *data_rows[start:end]])
                _report_progress(
                    progress_callback,
                    "split",
                    end,
                    generated_rows,
                    f"正在生成第 {index + 1}/{count} 个分表",
                )
        else:
            workbook = load_workbook(source, data_only=False)
            try:
                worksheet = _first_visible_sheet(workbook, source)
                last_row, width = _worksheet_bounds(worksheet, source)
                generated_rows = last_row - 1
                if generated_rows <= 0:
                    raise TableOperationError("总表只有表头，没有可拆分的数据。")
                count = (generated_rows + data_per_file - 1) // data_per_file
                output_paths = _split_output_paths(source, output_dir, count)
                _report_progress(
                    progress_callback, "split", 0, generated_rows, "准备生成分表……"
                )
                for index, path in enumerate(output_paths):
                    _check_cancel(cancel_check)
                    if status_callback:
                        status_callback(f"正在生成第 {index + 1}/{count} 个分表……")
                    start_row = 2 + index * data_per_file
                    end_row = min(last_row, start_row + data_per_file - 1)
                    _create_xlsx_output(
                        worksheet,
                        path,
                        [1, *range(start_row, end_row + 1)],
                        width,
                        completed_before=start_row - 2,
                        total_data_rows=generated_rows,
                        progress_message=f"正在生成第 {index + 1}/{count} 个分表",
                        progress_callback=progress_callback,
                        cancel_check=cancel_check,
                    )
                    _report_progress(
                        progress_callback,
                        "split",
                        end_row - 1,
                        generated_rows,
                        f"正在生成第 {index + 1}/{count} 个分表",
                    )
            finally:
                workbook.close()
        _check_cancel(cancel_check)
        _report_progress(progress_callback, "done", generated_rows, generated_rows, "拆分完成")
    except TableOperationCancelled:
        _remove_created_files(output_paths)
        raise

    return TableOperationResult(
        action=TABLE_ACTION_SPLIT,
        generated_rows=generated_rows,
        output_file_count=len(output_paths),
        source_file_count=1,
        output_dir=str(output_dir),
        output_path=str(output_paths[0]) if output_paths else None,
        output_paths=tuple(str(path) for path in output_paths),
    )


def _merge_input_files(folder: Path) -> tuple[str, list[Path]]:
    candidates = [
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in {".xlsx", ".csv"}
        and not path.name.startswith("~$")
        and not path.stem.startswith("合并总表_")
    ]
    if not candidates:
        raise TableOperationError("所选文件夹中没有可合并的 .xlsx 或 .csv 文件。")
    suffixes = {path.suffix.lower() for path in candidates}
    if len(suffixes) != 1:
        raise TableOperationError("同一次合并不能混合 .xlsx 和 .csv，请分别放入不同文件夹。")
    if len(candidates) < 2:
        raise TableOperationError("至少需要两个同格式分表才能合并。")
    return suffixes.pop(), sorted(candidates, key=lambda path: path.name.casefold())


def _header_mismatch_message(
    baseline: Sequence[object], current: Sequence[object], path: Path
) -> str:
    if len(baseline) != len(current):
        return (
            f"表头不一致：{path.name} 有 {len(current)} 列，"
            f"基准表有 {len(baseline)} 列。"
        )
    for index, (expected, actual) in enumerate(zip(baseline, current), start=1):
        if expected != actual:
            return (
                f"表头不一致：{path.name} 第 {index} 列应为“{expected}”，"
                f"实际为“{actual}”。"
            )
    return f"表头不一致：{path.name}。"


def _merge_csv(
    paths: Sequence[Path],
    output_path: Path,
    status_callback: Callable[[str], None] | None,
    progress_callback: ProgressCallback | None,
    cancel_check: CancelCheck | None,
) -> int:
    baseline: list[str] | None = None
    total_rows = 0
    _report_progress(progress_callback, "scan", 0, 0, "正在校验分表并统计数据量……")
    for index, path in enumerate(paths, start=1):
        _check_cancel(cancel_check)
        if status_callback:
            status_callback(f"正在校验第 {index}/{len(paths)} 个分表：{path.name}")
        _encoding, raw_rows = _decode_csv(path)
        header, data = _validate_csv_table(path, raw_rows)
        if baseline is None:
            baseline = header
        elif header != baseline:
            raise TableOperationError(_header_mismatch_message(baseline, header, path))
        total_rows += len(data)
    assert baseline is not None
    _report_progress(progress_callback, "merge", 0, total_rows, "准备合并分表……")

    temporary_path = _unique_path(
        output_path.with_name(f".{output_path.stem}_未完成{output_path.suffix}")
    )
    completed = 0
    try:
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            writer.writerow(baseline)
            for file_index, path in enumerate(paths, start=1):
                _check_cancel(cancel_check)
                if status_callback:
                    status_callback(
                        f"正在合并第 {file_index}/{len(paths)} 个分表：{path.name}"
                    )
                _encoding, raw_rows = _decode_csv(path)
                _header, data = _validate_csv_table(path, raw_rows)
                for start in range(0, len(data), 2_000):
                    _check_cancel(cancel_check)
                    chunk = data[start : start + 2_000]
                    writer.writerows(chunk)
                    completed += len(chunk)
                    _report_progress(
                        progress_callback,
                        "merge",
                        completed,
                        total_rows,
                        f"正在合并第 {file_index}/{len(paths)} 个分表",
                    )
        _check_cancel(cancel_check)
        _report_progress(
            progress_callback, "saving", completed, total_rows, "正在保存合并结果……"
        )
        os.replace(temporary_path, output_path)
        return completed
    except BaseException:
        _remove_created_files((temporary_path,))
        raise


@dataclass(frozen=True)
class _XlsxStreamInfo:
    path: Path
    title: str
    header: tuple[object, ...]
    width: int
    data_rows: int


def _reset_stream_dimensions(worksheet: object) -> None:
    reset_dimensions = getattr(worksheet, "reset_dimensions", None)
    if callable(reset_dimensions):
        reset_dimensions()


def _xlsx_stream_info(
    path: Path,
    file_index: int,
    file_count: int,
    progress_callback: ProgressCallback | None,
    cancel_check: CancelCheck | None,
) -> _XlsxStreamInfo:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = _first_visible_sheet(workbook, path)
        declared_rows = int(worksheet.max_row or 0)
        declared_columns = int(worksheet.max_column or 0)
        _reset_stream_dimensions(worksheet)
        rows = worksheet.iter_rows()
        try:
            header_cells = next(rows)
        except StopIteration as exc:
            raise TableOperationError(f"工作表为空，未找到表头：{path.name}") from exc
        header_values = tuple(cell.value for cell in header_cells)
        width = _last_nonempty_column(header_values)
        if width == 0:
            raise TableOperationError(f"第一行表头为空：{path.name}")
        header = header_values[:width]

        # 范围明显过小时不能相信 max_row；完整扫描真实 XML，既修复 A1/A1:C1，
        # 也为进度条取得准确总行数。正常文件直接使用声明行数，不增加预读耗时。
        suspicious_dimension = declared_rows <= 1 or declared_columns < width
        if suspicious_dimension:
            seen_rows = 0
            last_nonempty_row = 0
            for row in rows:
                seen_rows += 1
                values = tuple(cell.value for cell in row)
                if any(_is_nonempty(value) for value in values[width:]):
                    raise TableOperationError(
                        f"{path.name} 第 {seen_rows + 1} 行在表头范围之外还有数据，"
                        "请先补全表头。"
                    )
                if any(_is_nonempty(value) for value in values[:width]):
                    last_nonempty_row = seen_rows
                if seen_rows % 2_000 == 0:
                    _check_cancel(cancel_check)
                    _report_progress(
                        progress_callback,
                        "scan",
                        seen_rows,
                        0,
                        f"正在识别第 {file_index}/{file_count} 个分表的真实数据范围",
                    )
            data_rows = last_nonempty_row
        else:
            data_rows = max(0, declared_rows - 1)
        return _XlsxStreamInfo(path, worksheet.title, header, width, data_rows)
    finally:
        workbook.close()


def _copy_stream_header_cell(source: object, target_sheet: object) -> WriteOnlyCell:
    target = WriteOnlyCell(target_sheet, value=source.value)
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    return target


def _merge_xlsx(
    paths: Sequence[Path],
    output_path: Path,
    status_callback: Callable[[str], None] | None,
    progress_callback: ProgressCallback | None,
    cancel_check: CancelCheck | None,
) -> int:
    infos: list[_XlsxStreamInfo] = []
    baseline: tuple[object, ...] | None = None
    _report_progress(progress_callback, "scan", 0, 0, "正在校验分表并统计数据量……")
    for file_index, path in enumerate(paths, start=1):
        _check_cancel(cancel_check)
        if status_callback:
            status_callback(f"正在校验第 {file_index}/{len(paths)} 个分表：{path.name}")
        info = _xlsx_stream_info(
            path, file_index, len(paths), progress_callback, cancel_check
        )
        header = info.header
        if baseline is None:
            baseline = header
        elif header != baseline:
            raise TableOperationError(_header_mismatch_message(baseline, header, path))
        infos.append(info)

    if baseline is None or not infos:
        raise TableOperationError("没有读取到可合并的工作表。")

    total_rows = sum(info.data_rows for info in infos)
    _report_progress(progress_callback, "merge", 0, total_rows, "准备合并分表……")
    temporary_path = _unique_path(
        output_path.with_name(f".{output_path.stem}_未完成{output_path.suffix}")
    )
    output_workbook = Workbook(write_only=True)
    output_sheet = output_workbook.create_sheet(infos[0].title)
    generated_rows = 0
    try:
        for file_index, info in enumerate(infos, start=1):
            _check_cancel(cancel_check)
            if status_callback:
                status_callback(
                    f"正在合并第 {file_index}/{len(infos)} 个分表：{info.path.name}"
                )
            workbook = load_workbook(info.path, read_only=True, data_only=False)
            try:
                worksheet = _first_visible_sheet(workbook, info.path)
                _reset_stream_dimensions(worksheet)
                rows = worksheet.iter_rows()
                header_cells = next(rows)
                if file_index == 1:
                    output_sheet.append(
                        [
                            _copy_stream_header_cell(cell, output_sheet)
                            for cell in header_cells[: info.width]
                        ]
                    )

                pending_blank_rows = 0
                rows_since_report = 0
                source_row = 1
                for row in rows:
                    source_row += 1
                    values = tuple(cell.value for cell in row)
                    if any(_is_nonempty(value) for value in values[info.width:]):
                        raise TableOperationError(
                            f"{info.path.name} 第 {source_row} 行在表头范围之外还有数据，"
                            "请先补全表头。"
                        )
                    output_values = list(values[: info.width])
                    if len(output_values) < info.width:
                        output_values.extend([None] * (info.width - len(output_values)))
                    if any(_is_nonempty(value) for value in output_values):
                        for _ in range(pending_blank_rows):
                            output_sheet.append([None] * info.width)
                        generated_rows += pending_blank_rows
                        pending_blank_rows = 0
                        output_sheet.append(output_values)
                        generated_rows += 1
                    else:
                        pending_blank_rows += 1

                    rows_since_report += 1
                    if rows_since_report >= 2_000:
                        rows_since_report = 0
                        _check_cancel(cancel_check)
                        _report_progress(
                            progress_callback,
                            "merge",
                            generated_rows,
                            total_rows,
                            f"正在合并第 {file_index}/{len(infos)} 个分表",
                        )
            finally:
                workbook.close()
            _report_progress(
                progress_callback,
                "merge",
                generated_rows,
                total_rows,
                f"已读取第 {file_index}/{len(infos)} 个分表",
            )

        _check_cancel(cancel_check)
        _report_progress(
            progress_callback,
            "saving",
            generated_rows,
            max(total_rows, generated_rows),
            "正在保存合并结果，取消可能需要稍候……",
        )
        output_workbook.save(temporary_path)
        output_workbook.close()
        _check_cancel(cancel_check)
        os.replace(temporary_path, output_path)
        return generated_rows
    except BaseException:
        try:
            output_sheet.close()
        except Exception:
            pass
        output_workbook.close()
        _remove_created_files((temporary_path,))
        raise


def merge_tables(
    folder_path: Path,
    status_callback: Callable[[str], None] | None = None,
    now: datetime | None = None,
    progress_callback: ProgressCallback | None = None,
    cancel_check: CancelCheck | None = None,
) -> TableOperationResult:
    folder = folder_path.resolve()
    if not folder.is_dir():
        raise TableOperationError("请选择存在的分表文件夹。")
    suffix, paths = _merge_input_files(folder)
    output_dir = folder / OUTPUT_DIR_NAME
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = (now or datetime.now()).strftime("%Y%m%d_%H%M")
    output_path = _unique_path(
        output_dir / f"合并总表_{timestamp}{OUTPUT_SUFFIX}{suffix}"
    )
    _check_cancel(cancel_check)
    if suffix == ".csv":
        generated_rows = _merge_csv(
            paths, output_path, status_callback, progress_callback, cancel_check
        )
    else:
        generated_rows = _merge_xlsx(
            paths, output_path, status_callback, progress_callback, cancel_check
        )
    _report_progress(
        progress_callback, "done", generated_rows, generated_rows, "合并完成"
    )
    return TableOperationResult(
        action=TABLE_ACTION_MERGE,
        generated_rows=generated_rows,
        output_file_count=1,
        source_file_count=len(paths),
        output_dir=str(output_dir),
        output_path=str(output_path),
        output_paths=(str(output_path),),
    )
