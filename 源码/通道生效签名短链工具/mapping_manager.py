"""Manage the local channel/main-code workbook without requiring Excel."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook


class MappingError(Exception):
    pass


def _value(value: object) -> str:
    return str(value if value is not None else "").strip()


def _validate_sheet(sheet) -> None:
    expected = ("通道编号", "通道名称", "主码号", "省份")
    actual = tuple(_value(sheet.cell(1, column).value) for column in range(1, 5))
    if actual != expected:
        raise MappingError("主码号表第1行必须依次为：通道编号、通道名称、主码号、省份。")


def list_rows(path: Path) -> list[dict[str, str | int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        _validate_sheet(sheet)
        return [
            {
                "row": index,
                "channel": _value(sheet.cell(index, 1).value),
                "name": _value(sheet.cell(index, 2).value),
                "main": _value(sheet.cell(index, 3).value),
                "province": _value(sheet.cell(index, 4).value),
            }
            for index in range(2, sheet.max_row + 1)
            if any(sheet.cell(index, col).value not in (None, "") for col in range(1, 5))
        ]
    finally:
        workbook.close()


def change_row(
    path: Path, action: str, row: int, channel: str = "", name: str = "",
    main: str = "", province: str = "",
) -> None:
    if action not in {"add", "update", "delete"}:
        raise MappingError("不支持的主码号操作。")
    if action != "delete":
        if not re.fullmatch(r"\d{8}", main):
            raise MappingError("主码号必须是8位数字。")
        if not channel.strip() or not name.strip():
            raise MappingError("通道编号和通道名称不能为空。")
        if not any(operator in name for operator in ("移动", "联通", "电信")):
            raise MappingError("通道名称需包含“移动”“联通”或“电信”，以便正确匹配运营商。")
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        _validate_sheet(sheet)
        if action != "add" and not (2 <= row <= sheet.max_row):
            raise MappingError("所选记录已不存在，请刷新后重试。")
        if action != "delete":
            for index in range(2, sheet.max_row + 1):
                if index != row and _value(sheet.cell(index, 1).value) == channel and _value(sheet.cell(index, 3).value) == main:
                    raise MappingError("相同主码号和通道编号的记录已存在。")
        if action == "delete":
            sheet.delete_rows(row)
        else:
            target = sheet.max_row + 1 if action == "add" else row
            for col, value in enumerate((channel, name, main, province), 1):
                cell = sheet.cell(target, col, value)
                if col in (1, 3):
                    cell.number_format = "@"
        temp = path.with_name(f".{path.stem}.{uuid4().hex}.tmp.xlsx")
        try:
            workbook.save(temp)
            os.replace(temp, path)
        finally:
            temp.unlink(missing_ok=True)
    finally:
        workbook.close()


def run_cli(args: list[str]) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mapping-file", type=Path, required=True)
    parser.add_argument("--mapping-action", choices=("list", "add", "update", "delete"), required=True)
    parser.add_argument("--mapping-row", type=int, default=0)
    parser.add_argument("--mapping-channel", default="")
    parser.add_argument("--mapping-name", default="")
    parser.add_argument("--mapping-main", default="")
    parser.add_argument("--mapping-province", default="")
    values = parser.parse_args(args)
    try:
        if values.mapping_action == "list":
            print(json.dumps(list_rows(values.mapping_file), ensure_ascii=False))
        else:
            change_row(
                values.mapping_file, values.mapping_action, values.mapping_row,
                values.mapping_channel.strip(), values.mapping_name.strip(),
                values.mapping_main.strip(), values.mapping_province.strip(),
            )
            print(json.dumps({"ok": True}, ensure_ascii=False))
    except (MappingError, OSError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    return 0
