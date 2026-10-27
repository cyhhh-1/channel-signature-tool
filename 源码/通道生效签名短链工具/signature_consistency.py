# -*- coding: utf-8 -*-
"""签名一致性判定。

第一层：原签名字符必须按原顺序出现在企业名称中，允许跳过企业名称中的文字。
第二层：原签名任意位置必须包含地区库中的省、市、区、县或县级市名称，
并且该地区文字必须在企业名称中连续出现，不能把地区字符从不同位置拼接出来。
省级、地市级可使用简称；区、县、县级市只使用带后缀的完整名称。
企业名称中含“个体工商户”或“个人独资”的数据判定为失败。
可切换“只匹配地区”：忽略签名其余文字，只校验签名和企业名称是否含同一地区。
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REGION_LIBRARY_FILE = "地区库.xlsx"
SIGNATURE_HEADER = "签名"
ENTERPRISE_HEADER = "企业名称"
RESULT_HEADER = "判定结果"
RESULT_SUCCESS = "成功"
RESULT_FAILURE = "失败"
SIGNATURE_MATCH_FULL_RULE = "full_rule"
SIGNATURE_MATCH_REGION_ONLY = "region_only"
EXCLUDED_ENTERPRISE_TYPES = ("个体工商户", "个人独资")


class SignatureConsistencyError(Exception):
    """签名一致性判定的可读错误。"""


@dataclass(frozen=True)
class SignatureConsistencyResult:
    output_path: str
    source_rows: int
    first_layer_rows: int
    second_layer_rows: int
    first_layer_filtered_rows: int
    second_layer_filtered_rows: int
    enterprise_type_filtered_rows: int
    processed_sheets: int


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normal_header(value: object) -> str:
    return "".join(_text(value).split())


def _is_ordered_subsequence(signature: str, enterprise_name: str) -> bool:
    """签名字符按原顺序出现在企业名称中时返回 True，允许中间跳字。"""
    if not signature:
        return False
    signature_index = 0
    for character in enterprise_name:
        if character == signature[signature_index]:
            signature_index += 1
            if signature_index == len(signature):
                return True
    return False


class RegionMatcher:
    """按行政级别使用地区库名称检查签名中的地区文字。"""

    def __init__(self, token_priority: dict[str, int]) -> None:
        self.token_priority = token_priority
        self.token_lengths = sorted({len(token) for token in token_priority})

    @classmethod
    def from_workbook(cls, path: Path) -> "RegionMatcher":
        if not path.exists():
            raise SignatureConsistencyError(
                f"未找到地区库：{path.name}。请把它与程序放在同一文件夹。"
            )
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise SignatureConsistencyError(f"地区库无法读取：{exc}") from exc
        try:
            worksheet = workbook["地区库"] if "地区库" in workbook.sheetnames else workbook.worksheets[0]
            rows = worksheet.iter_rows(values_only=True)
            try:
                headers = [_normal_header(value) for value in next(rows)]
            except StopIteration as exc:
                raise SignatureConsistencyError("地区库为空。") from exc
            required = {"行政区级别", "完整名称", "可识别简称", "匹配优先级"}
            missing = required.difference(headers)
            if missing:
                raise SignatureConsistencyError(
                    "地区库缺少列：" + "、".join(sorted(missing))
                )
            full_index = headers.index("完整名称")
            short_index = headers.index("可识别简称")
            priority_index = headers.index("匹配优先级")
            level_index = headers.index("行政区级别")
            token_priority: dict[str, int] = {}
            for row in rows:
                try:
                    priority = int(row[priority_index])
                except (TypeError, ValueError, IndexError):
                    continue
                level = _text(row[level_index] if level_index < len(row) else "")
                # 省级、地市级允许简称；区县级简称容易与品牌词重名，
                # 因此区、县、县级市只允许带“区/县/市”等后缀的完整名称。
                allowed_indexes = (
                    (full_index, short_index)
                    if level in {"省级", "地市级"}
                    else (full_index,)
                )
                for index in allowed_indexes:
                    token = _text(row[index] if index < len(row) else "")
                    if len(token) < 2:
                        continue
                    token_priority[token] = max(priority, token_priority.get(token, 0))
            if not token_priority:
                raise SignatureConsistencyError("地区库中没有可用的完整名称或简称。")
            return cls(token_priority)
        finally:
            workbook.close()

    def has_region(self, signature: str, enterprise_name: str) -> bool:
        """签名和企业名称都连续包含同一合规地区文字时返回 True。"""
        if not signature or not enterprise_name:
            return False
        for start in range(len(signature)):
            for length in self.token_lengths:
                end = start + length
                if end > len(signature):
                    break
                if start == 0 and end == len(signature):
                    continue
                region = signature[start:end]
                if region in self.token_priority and region in enterprise_name:
                    return True
        return False


def _copy_cell(source: Cell, target: Cell, value: object | None = None) -> None:
    target.value = source.value if value is None else value
    if getattr(source, "has_style", False):
        target.font = copy(getattr(source, "font"))
        target.fill = copy(getattr(source, "fill"))
        target.border = copy(getattr(source, "border"))
        target.alignment = copy(getattr(source, "alignment"))
        target.protection = copy(getattr(source, "protection"))
    number_format = getattr(source, "number_format", None)
    if number_format:
        target.number_format = number_format
    hyperlink = getattr(source, "hyperlink", None)
    comment = getattr(source, "comment", None)
    if hyperlink:
        target._hyperlink = copy(hyperlink)
    if comment:
        target.comment = copy(comment)


def _find_header(worksheet, scan_rows: int = 50) -> tuple[int, int, int, int] | None:
    declared_max_row = getattr(worksheet, "max_row", None)
    max_scan_row = min(scan_rows, declared_max_row) if declared_max_row else scan_rows
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_scan_row),
        start=1,
    ):
        normalized = [_normal_header(cell.value) for cell in row]
        if SIGNATURE_HEADER in normalized and ENTERPRISE_HEADER in normalized:
            if RESULT_HEADER in normalized:
                raise SignatureConsistencyError(
                    f"工作表“{worksheet.title}”已包含“{RESULT_HEADER}”列，"
                    "请选择未经判定的原始文件。"
                )
            return (
                row_number,
                normalized.index(SIGNATURE_HEADER) + 1,
                normalized.index(ENTERPRISE_HEADER) + 1,
                len(row),
            )
    return None


def _unique_output_path(source: Path) -> Path:
    index = 0
    while True:
        suffix = "" if index == 0 else f"({index})"
        candidate = source.with_name(f"{source.stem}_输出结果{suffix}.xlsx")
        if not candidate.exists():
            return candidate
        index += 1


def _copy_column_widths(source_sheet, target_sheet, signature_column: int, max_column: int) -> None:
    source_dimensions = getattr(source_sheet, "column_dimensions", {})
    for source_column in range(1, max_column + 1):
        target_column = source_column if source_column <= signature_column else source_column + 1
        source_letter = get_column_letter(source_column)
        target_letter = get_column_letter(target_column)
        source_dimension = source_dimensions.get(source_letter)
        if source_dimension and source_dimension.width:
            target_sheet.column_dimensions[target_letter].width = source_dimension.width
    target_sheet.column_dimensions[get_column_letter(signature_column + 1)].width = 12


def _style_fallback_header(worksheet, row_number: int, max_column: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet.iter_cols(
        min_col=1,
        max_col=max_column,
        min_row=row_number,
        max_row=row_number,
    ):
        current = cell[0]
        if not current.has_style:
            current.fill = fill
            current.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
            current.alignment = Alignment(horizontal="center", vertical="center")


def process_signature_consistency(
    source_path: Path,
    region_library_path: Path,
    status_callback: Callable[[str], None] | None = None,
    match_mode: str = SIGNATURE_MATCH_FULL_RULE,
) -> SignatureConsistencyResult:
    source_path = source_path.resolve()
    if not source_path.exists() or not source_path.is_file():
        raise SignatureConsistencyError("请选择存在的 Excel 文件。")
    if source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SignatureConsistencyError("签名一致性判定只支持 .xlsx 和 .xlsm 文件。")
    if match_mode not in {SIGNATURE_MATCH_FULL_RULE, SIGNATURE_MATCH_REGION_ONLY}:
        raise SignatureConsistencyError("未知的签名判定规则。")
    if status_callback:
        status_callback("正在读取地区库……")
    matcher = RegionMatcher.from_workbook(region_library_path.resolve())
    try:
        source_workbook = load_workbook(source_path, read_only=True, data_only=False)
    except Exception as exc:
        raise SignatureConsistencyError(f"客户源文件无法读取：{exc}") from exc

    output_workbook = Workbook()
    output_workbook.remove(output_workbook.active)
    source_rows = 0
    first_layer_rows = 0
    second_layer_rows = 0
    enterprise_type_filtered_rows = 0
    processed_sheets = 0

    try:
        for worksheet in source_workbook.worksheets:
            reset_dimensions = getattr(worksheet, "reset_dimensions", None)
            if callable(reset_dimensions):
                reset_dimensions()
        for source_sheet in source_workbook.worksheets:
            header_info = _find_header(source_sheet)
            if header_info is None:
                continue
            processed_sheets += 1
            header_row, signature_column, enterprise_column, max_column = header_info
            target_sheet = output_workbook.create_sheet(source_sheet.title)
            source_view = getattr(source_sheet, "sheet_view", None)
            target_sheet.sheet_view.showGridLines = getattr(
                source_view, "showGridLines", True
            )
            _copy_column_widths(source_sheet, target_sheet, signature_column, max_column)

            # 保留表头及表头之前的说明行，在原“签名”后插入判定结果列。
            for target_row, source_row in enumerate(
                source_sheet.iter_rows(min_row=1, max_row=header_row),
                start=1,
            ):
                for source_column, source_cell in enumerate(
                    source_row[:max_column], start=1
                ):
                    target_column = (
                        source_column
                        if source_column <= signature_column
                        else source_column + 1
                    )
                    _copy_cell(source_cell, target_sheet.cell(target_row, target_column))
                if target_row == header_row:
                    _copy_cell(
                        source_row[signature_column - 1],
                        target_sheet.cell(target_row, signature_column + 1),
                        RESULT_HEADER,
                    )
                source_row_dimensions = getattr(source_sheet, "row_dimensions", {})
                row_dimension = source_row_dimensions.get(target_row)
                height = row_dimension.height if row_dimension else None
                if height:
                    target_sheet.row_dimensions[target_row].height = height

            target_data_row = header_row + 1
            for source_row_number, source_row in enumerate(
                source_sheet.iter_rows(min_row=header_row + 1),
                start=header_row + 1,
            ):
                signature = _text(source_row[signature_column - 1].value)
                enterprise_name = _text(source_row[enterprise_column - 1].value)
                has_data = any(cell.value is not None for cell in source_row[:max_column])
                result_text = ""
                if has_data:
                    source_rows += 1
                    result_text = RESULT_FAILURE
                    if any(keyword in enterprise_name for keyword in EXCLUDED_ENTERPRISE_TYPES):
                        enterprise_type_filtered_rows += 1
                    elif (
                        match_mode == SIGNATURE_MATCH_REGION_ONLY
                        and matcher.has_region(signature, enterprise_name)
                    ):
                        first_layer_rows += 1
                        second_layer_rows += 1
                        result_text = RESULT_SUCCESS
                    elif (
                        match_mode == SIGNATURE_MATCH_FULL_RULE
                        and _is_ordered_subsequence(signature, enterprise_name)
                    ):
                        first_layer_rows += 1
                        if matcher.has_region(signature, enterprise_name):
                            second_layer_rows += 1
                            result_text = RESULT_SUCCESS
                for source_column, source_cell in enumerate(
                    source_row[:max_column], start=1
                ):
                    target_column = (
                        source_column
                        if source_column <= signature_column
                        else source_column + 1
                    )
                    _copy_cell(source_cell, target_sheet.cell(target_data_row, target_column))
                if has_data:
                    _copy_cell(
                        source_row[signature_column - 1],
                        target_sheet.cell(target_data_row, signature_column + 1),
                        result_text,
                    )
                source_row_dimensions = getattr(source_sheet, "row_dimensions", {})
                row_dimension = source_row_dimensions.get(source_row_number)
                source_height = row_dimension.height if row_dimension else None
                if source_height:
                    target_sheet.row_dimensions[target_data_row].height = source_height
                target_data_row += 1

            last_column = get_column_letter(max_column + 1)
            target_sheet.freeze_panes = target_sheet.cell(header_row + 1, 1)
            target_sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(header_row, target_data_row - 1)}"
            _style_fallback_header(target_sheet, header_row, max_column + 1)
            if status_callback:
                status_callback(
                    f"已处理工作表“{source_sheet.title}”，当前成功 {second_layer_rows:,} 条……"
                )
    finally:
        source_workbook.close()

    if processed_sheets == 0:
        output_workbook.close()
        raise SignatureConsistencyError(
            "没有找到同时包含“签名”和“企业名称”列的工作表。"
        )

    output_path = _unique_output_path(source_path)
    temporary_path = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
    if status_callback:
        status_callback(
            f"正在保存 {source_rows:,} 条判定结果（成功 {second_layer_rows:,} 条）……"
        )
    try:
        output_workbook.save(temporary_path)
        temporary_path.replace(output_path)
    except Exception as exc:
        if temporary_path.exists():
            temporary_path.unlink()
        raise SignatureConsistencyError(f"输出结果保存失败：{exc}") from exc
    finally:
        output_workbook.close()

    return SignatureConsistencyResult(
        output_path=str(output_path),
        source_rows=source_rows,
        first_layer_rows=first_layer_rows,
        second_layer_rows=second_layer_rows,
        first_layer_filtered_rows=(
            source_rows - enterprise_type_filtered_rows - first_layer_rows
        ),
        second_layer_filtered_rows=first_layer_rows - second_layer_rows,
        enterprise_type_filtered_rows=enterprise_type_filtered_rows,
        processed_sheets=processed_sheets,
    )
