# -*- coding: utf-8 -*-
"""通道生效签名 / 生效短链 Excel 导入工具。

程序支持八种工作模式：
1. 生效短链：读取签名和引流内容，写入短链模板。
2. 生效签名：只读取签名，写入签名模板。
3. 混合表自动分流：所有签名去重后进入签名模板；“引流”非空的记录
   另外进入短链模板，并保证两个结果的子端口不重复。
4. 自家移动：从关键信息报备表取得签名和引流，按地区运营商
   文件夹中的签名总表文件名动态识别本公司主码号并查询业务类型。
5. 自家电信：汇总提交表和失败回执，剔除失败子端口后，
   使用子端口后缀匹配平台引流子码表，生成短链、签名、最终引流报备和签名错误原因结果。
6. 拆分/合并表：按包含表头的最大行数拆分 XLSX/CSV 总表，或合并文件夹内
   表头完全一致的同格式分表。
7. 签名一致性判定：可选“完整规则”或“只匹配地区”；签名和企业名称中
   必须连续包含同一合规地区名称，并在原表新增“判定结果”列。
8. 子端口分配：按主码号和签名编号对照最新历史总表，分别回填行业和营销空号。

安装版模板内置，主码号映射在每位用户的本地数据目录中管理。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import shutil
import sys
import threading
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from itertools import chain
from pathlib import Path
from typing import Callable, Iterable, Iterator, Sequence

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter

_TOOL_SOURCE_DIR = str(Path(__file__).resolve().parent)
if _TOOL_SOURCE_DIR not in sys.path:
    sys.path.insert(0, _TOOL_SOURCE_DIR)

from table_split_merge import (  # noqa: E402
    TABLE_ACTION_MERGE,
    TABLE_ACTION_SPLIT,
    TableOperationError,
    merge_tables,
    split_table,
)
from signature_consistency import (  # noqa: E402
    REGION_LIBRARY_FILE,
    SIGNATURE_MATCH_FULL_RULE,
    SIGNATURE_MATCH_REGION_ONLY,
    SignatureConsistencyError,
    process_signature_consistency,
)
from port_allocator import PortAllocationError, allocate_ports  # noqa: E402
from mapping_manager import run_cli as run_mapping_cli  # noqa: E402


APP_NAME = "通道生效签名短链程序"
SHORTLINK_TEMPLATE = "通道生效签名短链导入模板.xlsx"
SIGNATURE_TEMPLATE = "通道生效签名导入模板.xlsx"
FINAL_DIVERSION_TEMPLATE = "最终报备表模板.xlsx"
CHANNEL_MAPPING_FILE = "通道编号及对应主码号.xlsx"


def channel_mapping_path(data_dir: Path) -> Path:
    configured = os.environ.get("CHANNEL_TOOL_MAPPING_FILE", "").strip()
    return Path(configured).resolve() if configured else data_dir / CHANNEL_MAPPING_FILE

# 各模式使用稳定的英文内部值，界面上显示中文名称。
MODE_SHORTLINK = "shortlink"
MODE_SIGNATURE = "signature"
MODE_MIXED = "mixed"
MODE_SELF_MOBILE = "self_mobile"
MODE_SELF_TELECOM = "self_telecom"
MODE_TABLE_TOOLS = "table_tools"
MODE_SIGNATURE_CONSISTENCY = "signature_consistency"
MODE_PORT_ALLOCATION = "port_allocation"
# 保留旧常量名，避免已有测试或调用方失效；其含义现为“自家移动”。
MODE_SELF_OWNED = MODE_SELF_MOBILE
TELECOM_MANIFEST_FILE = "电信批次处理完成.json"
TELECOM_OUTPUT_DIR = "生成结果"

# 这些值表示客户明确写了“没有引流”，混合模式中按空白引流处理。
EMPTY_DIVERSION_PLACEHOLDERS = {
    "-",
    "—",
    "无",
    "暂无",
    "无引流",
    "不涉及",
    "n/a",
    "null",
}

SIGNATURE_HEADERS = {
    "签名",
    "短信签名",
    "客户签名",
    "报备签名",
    "签名名称",
    "供应商签名",
}
# 第三方生效签名文件有时不单独提供签名列，而是把签名放在模板内容开头。
# 该表头只供生效签名模式使用，避免影响短链和混合模式的表头判断。
SIGNATURE_TEMPLATE_CONTENT_HEADERS = {
    "模板内容",
}
# “引流”既可能包含链接，也可能包含号码，所以属于混合内容列。
COMBINED_HEADERS = {
    "引流",
    "引流内容",
    "引流信息",
    "短链及号码",
    "链接及号码",
    "引流号码和链接",
    "引流链接和号码",
}
LINK_HEADERS = {
    "链接",
    "短链",
    "引流链接",
    "引流域名",
    "短链地址",
    "链接地址",
    "域名",
    "网址",
    "url",
    "邮箱",
    "报备短链",
    "长链接",
}
NUMBER_HEADERS = {
    "号码",
    "号码1",
    "号码2",
    "引流号码",
    "引流号码1",
    "引流号码2",
    "电话号码",
    "电话",
    "手机号",
    "手机号码",
    "座机",
    "座机号码",
    "固话",
    "固定电话",
    "固定电话号码",
    "联系号码",
    "报备号码1",
    "报备号码2",
}

SELF_SOURCE_PORT_HEADERS = {"子端口号", "完整子端口", "完整子端口号", "真实发送端口号"}
SELF_SOURCE_PROVINCE_HEADERS = {"归属省", "省份", "接入省"}
SELF_SOURCE_STATUS_HEADERS = {"审核状态", "状态"}
SELF_SOURCE_MESSAGE_TYPE_HEADERS = {"消息类型"}
SELF_SOURCE_SHORTLINK_HEADERS = {"报备短链"}
SELF_SOURCE_NUMBER1_HEADERS = {"报备号码1"}
SELF_SOURCE_NUMBER2_HEADERS = {"报备号码2"}
SIGNATURE_TOTAL_PORT_HEADERS = {"子端口号"}
SIGNATURE_TOTAL_BUSINESS_HEADERS = {"业务类型"}
CHANNEL_CODE_HEADERS = {"通道编号"}
CHANNEL_NAME_HEADERS = {"通道名称"}
MAIN_CODE_HEADERS = {"主码号"}
TELECOM_SUBMISSION_MAIN_HEADERS = {"所属主端口"}
TELECOM_SUBMISSION_PORT_HEADERS = {"子端口号"}
TELECOM_SUBMISSION_BUSINESS_HEADERS = {"业务类型"}
TELECOM_SUBMISSION_SIGNATURE_HEADERS = {"子端口签名"}
TELECOM_SUBMISSION_ENTERPRISE_HEADERS = {"企业名称", "子端口企业名称"}
TELECOM_SUBMISSION_MESSAGE_HEADERS = {"短信模板", "内容"}
TELECOM_RECEIPT_MAIN_HEADERS = {"主端口号"}
TELECOM_RECEIPT_PORT_HEADERS = {"子端口号"}
TELECOM_RECEIPT_REASON_HEADERS = {"错误原因"}
TELECOM_PLATFORM_MAIN_HEADERS = {"主端口号"}
TELECOM_PLATFORM_PORT_HEADERS = {"子端口号"}
TELECOM_PLATFORM_SIGNATURE_HEADERS = {"企业签名"}
TELECOM_PLATFORM_LINK_HEADERS = {"加白链接"}
TELECOM_PLATFORM_NUMBER_HEADERS = {"加白号码"}
TELECOM_PLATFORM_STATUS_HEADERS = {"状态"}
CHANNEL_PROVINCE_HEADERS = {"省份"}

BUSINESS_TYPE_MAPPING = {
    "通知提醒": "行业",
    "广告促销": "营销",
    "行业": "行业",
    "营销": "营销",
}
MOBILE_VARIANT_REGULAR = "普通"
MOBILE_VARIANT_5G = "5G"
MOBILE_MESSAGE_TYPE_MAPPING = {
    "行短": MOBILE_VARIANT_REGULAR,
    "5G消息": MOBILE_VARIANT_5G,
}

SHORTLINK_OUTPUT_HEADERS = [
    "通道编号",
    "签名",
    "子端口",
    "运营商",
    "业务类型",
    "等级",
    "链接",
    "号码",
    "是否重保",
    "发送量",
    "错误信息",
    "备注",
]
SIGNATURE_OUTPUT_HEADERS = [
    "通道编号",
    "签名",
    "子端口",
    "运营商",
    "业务类型",
    "是否重保",
    "备注",
]
FINAL_DIVERSION_OPERATION_HEADER = "操作类型（1.新增修改2.删除）（注：填数字）"
FINAL_DIVERSION_OUTPUT_HEADERS = [
    "序号",
    "省份",
    "主端口号",
    "子端口号",
    "子端口企业名称",
    "号码1",
    "号码2",
    "链接",
    "用途说明",
    FINAL_DIVERSION_OPERATION_HEADER,
    "签名",
    "内容",
]
FINAL_DIVERSION_DEFAULT_PURPOSE = "用于发送通知申请白名单管理"
FINAL_DIVERSION_DEFAULT_OPERATION = 1
TELECOM_CATEGORY_MIXED = "mixed"
TELECOM_CATEGORY_DIVERSION = "diversion"
TELECOM_CATEGORY_SIGNATURE = "signature"
class ProcessingError(Exception):
    """可以直接显示给使用者的处理错误。"""


@dataclass(frozen=True)
class ErrorRecord:
    """保留异常数据的原始位置和原因，便于导出错误明细。"""
    source: str
    row_number: int
    signature: str
    raw_content: str
    reason: str
    category: str = ""


@dataclass(frozen=True)
class SelfOwnedErrorRecord:
    """自家通道模式的异常行，重点保留完整子端口和归属省。"""
    source: str
    row_number: int
    full_port: str
    province: str
    raw_content: str
    reason: str


@dataclass(frozen=True)
class SignatureTotalRecord:
    row_number: int
    full_port: str
    business_source: str
    raw_content: str


@dataclass(frozen=True)
class TelecomSubmissionRecord:
    source: str
    row_number: int
    main_code: str
    full_port: str
    signature: str
    business_type: str
    enterprise_name: str
    message_template: str
    category: str
    raw_content: str


@dataclass(frozen=True)
class TelecomReceiptFailureRecord:
    source: str
    row_number: int
    main_code: str
    full_port: str
    reason: str


@dataclass(frozen=True)
class TelecomFailureReasonRecord:
    main_code: str
    full_port: str
    signature: str
    reason: str
    status: str = "失败"


@dataclass(frozen=True)
class TelecomPlatformRecord:
    row_number: int
    suffix: str
    raw_link: str
    raw_number: str


@dataclass(frozen=True)
class ChannelMappingRecord:
    channel: str
    province: str


@dataclass(frozen=True)
class TelecomErrorRecord:
    source: str
    row_number: int
    main_code: str
    full_port: str
    signature: str
    raw_content: str
    reason: str


@dataclass(frozen=True)
class SelfOwnedOutputRecord:
    channel: str
    signature: str
    sub_port: str
    operator: str
    business_type: str
    important_support: str
    link: str = ""
    number: str = ""


@dataclass(frozen=True)
class FinalDiversionRecord:
    province: str
    main_code: str
    full_port: str
    enterprise_name: str
    number1: str
    number2: str
    link: str
    signature: str
    content: str


@dataclass(frozen=True)
class ProcessSummary:
    """一次运行的结果汇总；混合模式会同时返回两个输出路径。"""
    output_path: str | None
    error_path: str | None
    generated_rows: int
    duplicate_rows: int
    error_rows: int
    secondary_output_path: str | None = None
    shortlink_rows: int = 0
    signature_rows: int = 0
    shortlink_output_path: str | None = None
    signature_output_path: str | None = None
    skipped_rows: int = 0
    dirty_rows: int = 0
    failed_rows: int = 0
    batch_count: int = 0
    final_diversion_rows: int = 0
    final_diversion_output_path: str | None = None
    failure_reason_rows: int = 0
    failure_reason_output_path: str | None = None
    unrelated_failed_rows: int = 0
    table_action: str | None = None
    output_file_count: int = 0
    source_file_count: int = 0
    output_dir: str | None = None
    output_paths: tuple[str, ...] = ()
    source_rows: int = 0
    first_layer_rows: int = 0
    second_layer_rows: int = 0
    first_layer_filtered_rows: int = 0
    second_layer_filtered_rows: int = 0
    enterprise_type_filtered_rows: int = 0
    processed_sheets: int = 0


_RESULT_PATH_FIELDS = (
    "output_path",
    "secondary_output_path",
    "error_path",
    "shortlink_output_path",
    "signature_output_path",
    "final_diversion_output_path",
    "failure_reason_output_path",
)

_GROUPED_OUTPUT_MODE_LABELS = {
    MODE_SELF_MOBILE: "自家移动",
    MODE_SELF_TELECOM: "自家电信",
    MODE_PORT_ALLOCATION: "子端口分配",
    MODE_MIXED: "混合表自动分流",
}


def _available_destination(path: Path, reserved: set[Path]) -> Path:
    """返回不覆盖现有结果文件、也不与本批次其他文件冲突的路径。"""
    candidate = path
    index = 1
    while candidate.exists() or candidate in reserved:
        candidate = path.with_name(f"{path.stem}({index}){path.suffix}")
        index += 1
    reserved.add(candidate)
    return candidate


def _grouped_result_directory(root: Path, mode: str) -> Path:
    label = _GROUPED_OUTPUT_MODE_LABELS[mode]
    folder_name = f"生成结果_{label}_{datetime.now():%m%d-%H时%M分}"
    candidate = root / folder_name
    index = 1
    while candidate.exists():
        candidate = root / f"{folder_name}({index})"
        index += 1
    return candidate


def relocate_summary_outputs(
    summary: ProcessSummary,
    destination_dir: Path | None,
    mode: str,
) -> ProcessSummary:
    """把本次新生成的结果归档到用户指定目录，并同步返回的新路径。"""
    if destination_dir is None:
        return summary
    selected_root = destination_dir.expanduser().resolve()
    destination = (
        _grouped_result_directory(selected_root, mode)
        if mode in _GROUPED_OUTPUT_MODE_LABELS
        else selected_root
    )
    try:
        destination.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ProcessingError(f"无法创建另存目录：{destination}") from exc

    values: list[str] = []
    for field_name in _RESULT_PATH_FIELDS:
        value = getattr(summary, field_name)
        if value:
            values.append(value)
    values.extend(summary.output_paths)

    relocated: dict[str, str] = {}
    reserved: set[Path] = set()
    source_parents: set[Path] = set()
    for value in values:
        source = Path(value).resolve()
        source_key = os.path.normcase(str(source))
        if source_key in relocated:
            continue
        if not source.is_file():
            continue
        if source.parent == destination:
            relocated[source_key] = str(source)
            reserved.add(source)
            continue
        target = _available_destination(destination / source.name, reserved)
        try:
            shutil.move(str(source), str(target))
        except OSError as exc:
            raise ProcessingError(f"无法将结果另存到：{target}") from exc
        relocated[source_key] = str(target)
        source_parents.add(source.parent)

    # 部分业务会先在源文件附近创建默认“生成结果”目录。文件归档后只清理
    # 已经完全为空的目录，绝不删除仍含输入文件或旧结果的目录。
    for parent in sorted(source_parents, key=lambda value: len(value.parts), reverse=True):
        if parent == destination or parent == selected_root:
            continue
        try:
            parent.rmdir()
        except OSError:
            pass

    def moved(value: str | None) -> str | None:
        if not value:
            return value
        key = os.path.normcase(str(Path(value).resolve()))
        return relocated.get(key, value)

    updates = {field_name: moved(getattr(summary, field_name)) for field_name in _RESULT_PATH_FIELDS}
    updates["output_paths"] = tuple(moved(value) or value for value in summary.output_paths)
    updates["output_dir"] = str(destination)
    return replace(summary, **updates)


def application_dir() -> Path:
    """返回程序所在目录；打包后使用 EXE 所在目录寻找外置模板。"""
    # 便携启动器会把自身目录放入环境变量，使内置 Python 仍能找到 EXE 旁的模板。
    launcher_dir = os.environ.get("CHANNEL_TOOL_HOME", "").strip()
    if launcher_dir:
        return Path(launcher_dir).resolve()
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def normalize_header(value: object) -> str:
    text = value_to_text(value).lower()
    return re.sub(r"[\s\[\]【】()（）:：_\-]+", "", text)


def value_to_text(value: object, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, int):
        text = str(value)
        if number_format and re.fullmatch(r"0+", number_format):
            text = text.zfill(len(number_format))
        return text
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            text = str(int(value))
            if number_format and re.fullmatch(r"0+", number_format):
                text = text.zfill(len(number_format))
            return text
        return format(value, ".15g")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


NORMALIZED_SIGNATURE_HEADERS = {normalize_header(value) for value in SIGNATURE_HEADERS}
NORMALIZED_SIGNATURE_TEMPLATE_CONTENT_HEADERS = {
    normalize_header(value) for value in SIGNATURE_TEMPLATE_CONTENT_HEADERS
}
NORMALIZED_COMBINED_HEADERS = {normalize_header(value) for value in COMBINED_HEADERS}
NORMALIZED_LINK_HEADERS = {normalize_header(value) for value in LINK_HEADERS}
NORMALIZED_NUMBER_HEADERS = {normalize_header(value) for value in NUMBER_HEADERS}
NORMALIZED_SELF_SOURCE_PORT_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_PORT_HEADERS
}
NORMALIZED_SELF_SOURCE_PROVINCE_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_PROVINCE_HEADERS
}
NORMALIZED_SELF_SOURCE_STATUS_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_STATUS_HEADERS
}
NORMALIZED_SELF_SOURCE_MESSAGE_TYPE_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_MESSAGE_TYPE_HEADERS
}
NORMALIZED_SELF_SOURCE_SHORTLINK_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_SHORTLINK_HEADERS
}
NORMALIZED_SELF_SOURCE_NUMBER1_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_NUMBER1_HEADERS
}
NORMALIZED_SELF_SOURCE_NUMBER2_HEADERS = {
    normalize_header(value) for value in SELF_SOURCE_NUMBER2_HEADERS
}
NORMALIZED_SIGNATURE_TOTAL_PORT_HEADERS = {
    normalize_header(value) for value in SIGNATURE_TOTAL_PORT_HEADERS
}
NORMALIZED_SIGNATURE_TOTAL_BUSINESS_HEADERS = {
    normalize_header(value) for value in SIGNATURE_TOTAL_BUSINESS_HEADERS
}
NORMALIZED_CHANNEL_CODE_HEADERS = {normalize_header(value) for value in CHANNEL_CODE_HEADERS}
NORMALIZED_CHANNEL_NAME_HEADERS = {normalize_header(value) for value in CHANNEL_NAME_HEADERS}
NORMALIZED_MAIN_CODE_HEADERS = {normalize_header(value) for value in MAIN_CODE_HEADERS}
NORMALIZED_TELECOM_SUBMISSION_MAIN_HEADERS = {
    normalize_header(value) for value in TELECOM_SUBMISSION_MAIN_HEADERS
}
NORMALIZED_TELECOM_SUBMISSION_PORT_HEADERS = {
    normalize_header(value) for value in TELECOM_SUBMISSION_PORT_HEADERS
}
NORMALIZED_TELECOM_SUBMISSION_BUSINESS_HEADERS = {
    normalize_header(value) for value in TELECOM_SUBMISSION_BUSINESS_HEADERS
}
NORMALIZED_TELECOM_SUBMISSION_SIGNATURE_HEADERS = {
    normalize_header(value) for value in TELECOM_SUBMISSION_SIGNATURE_HEADERS
}
NORMALIZED_TELECOM_SUBMISSION_ENTERPRISE_HEADERS = {
    normalize_header(value) for value in TELECOM_SUBMISSION_ENTERPRISE_HEADERS
}
NORMALIZED_TELECOM_SUBMISSION_MESSAGE_HEADERS = {
    normalize_header(value) for value in TELECOM_SUBMISSION_MESSAGE_HEADERS
}
NORMALIZED_TELECOM_RECEIPT_MAIN_HEADERS = {
    normalize_header(value) for value in TELECOM_RECEIPT_MAIN_HEADERS
}
NORMALIZED_TELECOM_RECEIPT_PORT_HEADERS = {
    normalize_header(value) for value in TELECOM_RECEIPT_PORT_HEADERS
}
NORMALIZED_TELECOM_RECEIPT_REASON_HEADERS = {
    normalize_header(value) for value in TELECOM_RECEIPT_REASON_HEADERS
}
NORMALIZED_TELECOM_PLATFORM_MAIN_HEADERS = {
    normalize_header(value) for value in TELECOM_PLATFORM_MAIN_HEADERS
}
NORMALIZED_TELECOM_PLATFORM_PORT_HEADERS = {
    normalize_header(value) for value in TELECOM_PLATFORM_PORT_HEADERS
}
NORMALIZED_TELECOM_PLATFORM_SIGNATURE_HEADERS = {
    normalize_header(value) for value in TELECOM_PLATFORM_SIGNATURE_HEADERS
}
NORMALIZED_TELECOM_PLATFORM_LINK_HEADERS = {
    normalize_header(value) for value in TELECOM_PLATFORM_LINK_HEADERS
}
NORMALIZED_TELECOM_PLATFORM_NUMBER_HEADERS = {
    normalize_header(value) for value in TELECOM_PLATFORM_NUMBER_HEADERS
}
NORMALIZED_TELECOM_PLATFORM_STATUS_HEADERS = {
    normalize_header(value) for value in TELECOM_PLATFORM_STATUS_HEADERS
}
NORMALIZED_CHANNEL_PROVINCE_HEADERS = {
    normalize_header(value) for value in CHANNEL_PROVINCE_HEADERS
}


def decode_text_file(path: Path) -> str:
    """兼容 UTF-8、UTF-16、GB18030 和 Big5 等常见中文编码。"""
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    encodings = ("utf-8-sig", "gb18030", "big5")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ProcessingError(f"无法识别文件编码：{path.name}。请另存为 UTF-8 或 GBK 后重试。")


def xlsx_row_values(cells: Sequence[Cell]) -> list[str]:
    return [value_to_text(cell.value, cell.number_format) for cell in cells]


def reset_input_workbook_dimensions(workbook: object) -> None:
    """只读输入表一律忽略 XLSX 声明范围，按工作表 XML 的真实单元格读取。

    部分平台或运营商导出的工作表会把实际数万行数据错误声明成 A1、A1:C1
    等较小范围。ReadOnlyWorksheet.reset_dimensions() 只重置内存中的读取边界，
    不扫描、不保存也不修改源文件内容。
    """
    for worksheet in getattr(workbook, "worksheets", []):
        reset_dimensions = getattr(worksheet, "reset_dimensions", None)
        if callable(reset_dimensions):
            reset_dimensions()


def iter_xlsx_sheets(path: Path) -> Iterator[tuple[str, Iterator[tuple[int, list[str]]]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ProcessingError(f"无法读取 Excel 文件“{path.name}”：{exc}") from exc

    try:
        reset_input_workbook_dimensions(workbook)
        for worksheet in workbook.worksheets:
            rows = (
                (row_number, xlsx_row_values(cells))
                for row_number, cells in enumerate(worksheet.iter_rows(), start=1)
            )
            yield worksheet.title, rows
    finally:
        workbook.close()


def iter_csv_sheets(path: Path) -> Iterator[tuple[str, Iterator[tuple[int, list[str]]]]]:
    text = decode_text_file(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows = ((row_number, [value.strip() for value in row]) for row_number, row in enumerate(reader, 1))
    yield path.stem, rows


def iter_tabular_sheets(path: Path) -> Iterator[tuple[str, Iterator[tuple[int, list[str]]]]]:
    """把 Excel/CSV 统一转换成“工作表名、行号、文本值”的迭代器。"""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        yield from iter_xlsx_sheets(path)
    elif suffix == ".csv":
        yield from iter_csv_sheets(path)
    elif suffix == ".xls":
        raise ProcessingError("不再支持旧版 .xls，请先在 Excel 中另存为 .xlsx 后再选择。")
    else:
        raise ProcessingError(f"不支持的文件格式：{suffix or '无扩展名'}")


def normalize_signature(value: object) -> str:
    """去除已有中英文方括号后，统一添加一层中文【】。"""
    text = value_to_text(value)
    if not text:
        return ""
    bracket_pairs = (("【", "】"), ("[", "]"), ("［", "］"))
    changed = True
    while changed and len(text) >= 2:
        changed = False
        for left, right in bracket_pairs:
            if text.startswith(left) and text.endswith(right):
                text = text[len(left) : -len(right)].strip()
                changed = True
                break
    return f"【{text}】" if text else ""


def signature_without_brackets(value: object) -> str:
    """返回不带最外层方括号的签名，供最终引流报备表使用。"""
    normalized = normalize_signature(value)
    return normalized[1:-1] if normalized.startswith("【") and normalized.endswith("】") else normalized


def extract_leading_signature(value: object) -> str:
    """从模板内容开头提取一层中英文方括号签名。"""
    text = value_to_text(value)
    match = re.match(r"^(【[^】]+】|\[[^\]]+\]|［[^］]+］)", text)
    return normalize_signature(match.group(1)) if match else ""


def is_signature_title(value: str) -> bool:
    compact = re.sub(r"\s+", "", value)
    return bool(
        re.fullmatch(r"(?:中国)?(?:移动|联通|电信|三网)?(?:短信)?签名[:：]?", compact)
        or normalize_header(compact) in NORMALIZED_SIGNATURE_HEADERS
    )


def split_items(value: object) -> list[str]:
    """拆分多个引流内容；斜杠不是分隔符，网址路径会完整保留。"""
    text = value_to_text(value)
    if not text:
        return []
    return [
        item.strip()
        for item in re.split(r"\s*(?:&&|\||｜|,|，|;|；|\r?\n)\s*", text)
        if item.strip()
    ]


def is_number_item(value: str) -> bool:
    """识别普通号码、区号号码、400 号码及带国家码的号码。"""
    text = value.strip()
    if not re.fullmatch(r"\+?\d[\d\-()（）\s]*\d|\d", text):
        return False
    digits = re.sub(r"\D", "", text)
    return len(digits) >= 3


def ordered_unique(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def raw_row_text(row: Sequence[str], max_chars: int = 1000) -> str:
    text = " | ".join(value for value in row if value)
    return text if len(text) <= max_chars else f"{text[: max_chars - 3]}..."


def find_signature_header(buffer: Sequence[tuple[int, list[str]]]) -> tuple[int, int] | None:
    for buffer_index, (_, row) in enumerate(buffer):
        for column_index, value in enumerate(row):
            if normalize_header(value) in NORMALIZED_SIGNATURE_HEADERS:
                return buffer_index, column_index
    return None


def find_signature_template_content_header(
    buffer: Sequence[tuple[int, list[str]]],
) -> tuple[int, int] | None:
    for buffer_index, (_, row) in enumerate(buffer):
        for column_index, value in enumerate(row):
            if normalize_header(value) in NORMALIZED_SIGNATURE_TEMPLATE_CONTENT_HEADERS:
                return buffer_index, column_index
    return None


@dataclass(frozen=True)
class ShortlinkColumns:
    signature: int
    item_columns: tuple[tuple[int, str], ...]


@dataclass(frozen=True)
class MixedColumns:
    signature: int
    item_columns: tuple[tuple[int, str], ...]


def find_shortlink_header(buffer: Sequence[tuple[int, list[str]]]) -> tuple[int, ShortlinkColumns] | None:
    for buffer_index, (_, row) in enumerate(buffer):
        signature_column: int | None = None
        item_columns: list[tuple[int, str]] = []
        for column_index, value in enumerate(row):
            header = normalize_header(value)
            if header in NORMALIZED_SIGNATURE_HEADERS and signature_column is None:
                signature_column = column_index
            elif header in NORMALIZED_COMBINED_HEADERS:
                item_columns.append((column_index, "combined"))
            elif header in NORMALIZED_LINK_HEADERS:
                item_columns.append((column_index, "link"))
            elif header in NORMALIZED_NUMBER_HEADERS:
                item_columns.append((column_index, "number"))
        if signature_column is not None and item_columns:
            return buffer_index, ShortlinkColumns(signature_column, tuple(item_columns))
    return None


def find_mixed_header(buffer: Sequence[tuple[int, list[str]]]) -> tuple[int, MixedColumns] | None:
    """混合模式接受综合引流列，也接受分开的链接列和号码列。"""
    header = find_shortlink_header(buffer)
    if header is None:
        return None
    buffer_index, columns = header
    return buffer_index, MixedColumns(columns.signature, columns.item_columns)


def is_blank_diversion(value: object) -> bool:
    """空字符串、纯空格和常见“无引流”占位词都视为没有引流。"""
    text = value_to_text(value)
    return not text or text.lower() in EMPTY_DIVERSION_PLACEHOLDERS


def extract_shortlink_groups(
    signature: str,
    row: Sequence[str],
    columns: ShortlinkColumns,
) -> list[tuple[str, str, str]]:
    """把一行引流拆成每组最多 1 个链接和 2 个号码。"""
    links: list[str] = []
    numbers: list[str] = []
    for column_index, _kind in sorted(columns.item_columns):
        value = row[column_index] if column_index < len(row) else ""
        if is_blank_diversion(value):
            continue
        for item in split_items(value):
            if is_number_item(item):
                numbers.append(item)
            else:
                links.append(item)
    links = ordered_unique(links)
    numbers = ordered_unique(numbers)

    if not links and not numbers:
        return []

    # 链接逐个分配；号码每两个一组，超出的内容生成新的子端口行。
    group_count = max(len(links), math.ceil(len(numbers) / 2), 1)
    groups: list[tuple[str, str, str]] = []
    for group_index in range(group_count):
        link = links[group_index] if group_index < len(links) else ""
        number = ",".join(numbers[group_index * 2 : group_index * 2 + 2])
        groups.append((signature, link, number))
    return groups


def buffer_rows(
    rows: Iterator[tuple[int, list[str]]], limit: int = 20
) -> tuple[list[tuple[int, list[str]]], Iterator[tuple[int, list[str]]]]:
    buffered: list[tuple[int, list[str]]] = []
    for _ in range(limit):
        try:
            buffered.append(next(rows))
        except StopIteration:
            break
    return buffered, rows


def read_signature_records(path: Path) -> tuple[list[tuple[str]], list[ErrorRecord], int]:
    """读取、规范化并去重生效签名数据。"""
    records: list[tuple[str]] = []
    errors: list[ErrorRecord] = []
    seen: set[str] = set()
    duplicate_count = 0

    def accept(source: str, row_number: int, raw_signature: str, raw: str) -> None:
        nonlocal duplicate_count
        if not raw_signature:
            if raw:
                errors.append(ErrorRecord(source, row_number, "", raw, "缺少签名"))
            return
        if is_signature_title(raw_signature):
            return
        signature = normalize_signature(raw_signature)
        if not signature:
            errors.append(ErrorRecord(source, row_number, raw_signature, raw, "签名为空或格式无效"))
            return
        if signature in seen:
            duplicate_count += 1
            return
        seen.add(signature)
        records.append((signature,))

    if path.suffix.lower() == ".txt":
        for row_number, line in enumerate(decode_text_file(path).splitlines(), start=1):
            text = line.strip()
            if not text:
                continue
            accept(path.name, row_number, text, text)
        return records, errors, duplicate_count

    for sheet_name, row_iterator in iter_tabular_sheets(path):
        iterator = iter(row_iterator)
        buffered, remaining = buffer_rows(iterator)
        header = find_signature_header(buffered)
        template_content_mode = False
        if header is None:
            header = find_signature_template_content_header(buffered)
            template_content_mode = header is not None
        if header is None:
            signature_column = 0
            data_rows = chain(buffered, remaining)
        else:
            header_index, signature_column = header
            data_rows = chain(buffered[header_index + 1 :], remaining)

        for row_number, row in data_rows:
            if not any(row):
                continue
            raw_signature = row[signature_column] if signature_column < len(row) else ""
            raw = raw_row_text(row)
            if template_content_mode:
                signature = extract_leading_signature(raw_signature)
                if not signature:
                    errors.append(
                        ErrorRecord(
                            sheet_name,
                            row_number,
                            "",
                            raw,
                            "模板内容没有以【签名】开头",
                        )
                    )
                    continue
                accept(sheet_name, row_number, signature, raw)
            else:
                accept(sheet_name, row_number, raw_signature, raw)

    return records, errors, duplicate_count


def read_shortlink_records(
    path: Path,
) -> tuple[list[tuple[str, str, str]], list[ErrorRecord], int]:
    """读取生效短链数据，并按模板容量拆分、精确去重。"""
    records: list[tuple[str, str, str]] = []
    errors: list[ErrorRecord] = []
    seen: set[tuple[str, str, str]] = set()
    duplicate_count = 0
    sheets_with_headers = 0

    for sheet_name, row_iterator in iter_tabular_sheets(path):
        iterator = iter(row_iterator)
        buffered, remaining = buffer_rows(iterator)
        header = find_shortlink_header(buffered)
        if header is None:
            continue
        sheets_with_headers += 1
        header_index, columns = header
        data_rows = chain(buffered[header_index + 1 :], remaining)

        for row_number, row in data_rows:
            if not any(row):
                continue
            raw_signature = row[columns.signature] if columns.signature < len(row) else ""
            raw = raw_row_text(row)
            if not raw_signature:
                errors.append(ErrorRecord(sheet_name, row_number, "", raw, "缺少签名"))
                continue
            if normalize_header(raw_signature) in NORMALIZED_SIGNATURE_HEADERS:
                repeated_headers = {
                    normalize_header(row[index])
                    for index, _kind in columns.item_columns
                    if index < len(row)
                }
                known_item_headers = (
                    NORMALIZED_COMBINED_HEADERS
                    | NORMALIZED_LINK_HEADERS
                    | NORMALIZED_NUMBER_HEADERS
                )
                if repeated_headers & known_item_headers:
                    continue
            signature = normalize_signature(raw_signature)
            if not signature:
                errors.append(
                    ErrorRecord(sheet_name, row_number, raw_signature, raw, "签名为空或格式无效")
                )
                continue

            groups = extract_shortlink_groups(signature, row, columns)
            if not groups:
                errors.append(
                    ErrorRecord(sheet_name, row_number, signature, raw, "没有可识别的链接或号码")
                )
                continue
            for key in groups:
                if key in seen:
                    duplicate_count += 1
                    continue
                seen.add(key)
                records.append(key)

    if sheets_with_headers == 0:
        raise ProcessingError(
            "没有找到生效短链数据表头。文件中必须包含“签名”，并至少包含“引流/链接/短链/号码”之一。"
        )
    return records, errors, duplicate_count


def read_mixed_records(
    path: Path,
) -> tuple[
    list[tuple[str, str, str]],
    list[tuple[str]],
    list[ErrorRecord],
    int,
    int,
]:
    """读取混合表：所有签名去重导出，有引流的记录另外导出短链。"""
    shortlink_records: list[tuple[str, str, str]] = []
    signature_records: list[tuple[str]] = []
    errors: list[ErrorRecord] = []
    seen_shortlinks: set[tuple[str, str, str]] = set()
    seen_signatures: set[str] = set()
    shortlink_duplicate_count = 0
    signature_duplicate_count = 0
    sheets_with_headers = 0

    for sheet_name, row_iterator in iter_tabular_sheets(path):
        iterator = iter(row_iterator)
        buffered, remaining = buffer_rows(iterator)
        header = find_mixed_header(buffered)
        if header is None:
            continue
        sheets_with_headers += 1
        header_index, columns = header
        data_rows = chain(buffered[header_index + 1 :], remaining)
        shortlink_columns = ShortlinkColumns(
            columns.signature,
            columns.item_columns,
        )

        for row_number, row in data_rows:
            if not any(row):
                continue
            raw_signature = row[columns.signature] if columns.signature < len(row) else ""
            raw_diversions = [
                row[column_index] if column_index < len(row) else ""
                for column_index, _kind in columns.item_columns
            ]
            raw = raw_row_text(row)
            # 只依据已识别的引流、链接和号码列，不根据短信模板等说明列猜测。
            category = (
                "生效签名"
                if all(is_blank_diversion(value) for value in raw_diversions)
                else "生效短链"
            )

            if (
                normalize_header(raw_signature) in NORMALIZED_SIGNATURE_HEADERS
                and any(
                    normalize_header(value)
                    in (
                        NORMALIZED_COMBINED_HEADERS
                        | NORMALIZED_LINK_HEADERS
                        | NORMALIZED_NUMBER_HEADERS
                    )
                    for value in raw_diversions
                )
            ):
                continue
            if not raw_signature:
                errors.append(
                    ErrorRecord(sheet_name, row_number, "", raw, "缺少签名", category)
                )
                continue
            signature = normalize_signature(raw_signature)
            if not signature:
                errors.append(
                    ErrorRecord(
                        sheet_name,
                        row_number,
                        raw_signature,
                        raw,
                        "签名为空或格式无效",
                        category,
                    )
                )
                continue

            # 签名模板覆盖所有有效签名，不受该行是否包含号码或链接影响。
            # 签名重复时仍需继续处理短链，以免丢失同签名下的不同引流内容。
            if signature in seen_signatures:
                signature_duplicate_count += 1
            else:
                seen_signatures.add(signature)
                signature_records.append((signature,))

            if category == "生效签名":
                continue

            groups = extract_shortlink_groups(signature, row, shortlink_columns)
            if not groups:
                errors.append(
                    ErrorRecord(
                        sheet_name,
                        row_number,
                        signature,
                        raw,
                        "没有可识别的链接或号码",
                        "生效短链",
                    )
                )
                continue
            for key in groups:
                if key in seen_shortlinks:
                    shortlink_duplicate_count += 1
                    continue
                seen_shortlinks.add(key)
                shortlink_records.append(key)

    if sheets_with_headers == 0:
        raise ProcessingError(
            "没有找到混合表表头。文件中必须包含“签名”，并至少包含“引流/链接/短链/号码”之一。"
        )
    return (
        shortlink_records,
        signature_records,
        errors,
        shortlink_duplicate_count,
        signature_duplicate_count,
    )


def read_first_xlsx_sheet(path: Path) -> tuple[str, list[tuple[int, list[str]]]]:
    """只读取 Excel 的第一个工作表；自家通道的所有业务表都遵循此规则。"""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ProcessingError(f"无法读取 Excel 文件“{path.name}”：{exc}") from exc
    try:
        if not workbook.worksheets:
            raise ProcessingError(f"Excel 文件“{path.name}”没有工作表。")
        reset_input_workbook_dimensions(workbook)
        worksheet = workbook.worksheets[0]
        rows = [
            (row_number, xlsx_row_values(cells))
            for row_number, cells in enumerate(worksheet.iter_rows(), start=1)
        ]
        return worksheet.title, rows
    finally:
        workbook.close()


def find_required_header(
    rows: Sequence[tuple[int, list[str]]],
    required: dict[str, set[str]],
    description: str,
) -> tuple[int, dict[str, int]]:
    """在前 20 行中寻找同时包含全部必需字段的表头。"""
    for buffer_index, (_row_number, row) in enumerate(rows[:20]):
        normalized = [normalize_header(value) for value in row]
        columns: dict[str, int] = {}
        for key, aliases in required.items():
            for column_index, header in enumerate(normalized):
                if header in aliases:
                    columns[key] = column_index
                    break
        if len(columns) == len(required):
            return buffer_index, columns
    labels = "、".join(required)
    raise ProcessingError(f"{description}表头不正确，必须包含：{labels}。")


def find_optional_column(row: Sequence[str], aliases: set[str]) -> int | None:
    """在已确定的表头行中查找可选列。"""
    for column_index, value in enumerate(row):
        if normalize_header(value) in aliases:
            return column_index
    return None


def normalize_region_name(value: object) -> str:
    """把“甘肃省”等常见写法转换为总表文件名使用的“甘肃”。"""
    text = value_to_text(value)
    for suffix in ("壮族自治区", "回族自治区", "维吾尔自治区", "自治区", "特别行政区", "省", "市"):
        if text.endswith(suffix) and len(text) > len(suffix):
            return text[: -len(suffix)]
    return text


def normalize_full_port(value: object) -> str:
    return re.sub(r"\s+", "", value_to_text(value))


def map_business_type(value: object) -> str:
    text = re.sub(r"\s+", "", value_to_text(value))
    return BUSINESS_TYPE_MAPPING.get(text, "")


def map_mobile_message_type(value: object) -> str:
    """只接受运营商当前明确约定的普通移动和 5G 消息类型。"""
    text = re.sub(r"\s+", "", value_to_text(value))
    return MOBILE_MESSAGE_TYPE_MAPPING.get(text, "")


def classify_mobile_channel(channel_name: object) -> str:
    """从现有不统一的通道名称中区分普通移动和 5G 移动。"""
    compact = normalize_header(channel_name).upper().replace("５", "5").replace("Ｇ", "G")
    return MOBILE_VARIANT_5G if "5G" in compact else MOBILE_VARIANT_REGULAR


def load_channel_mapping(path: Path) -> dict[tuple[str, str, str], list[str]]:
    """读取主码号、运营商及移动类型对应通道；非 8 位主码号忽略。"""
    _sheet_name, rows = read_first_xlsx_sheet(path)
    header_index, columns = find_required_header(
        rows,
        {
            "通道编号": NORMALIZED_CHANNEL_CODE_HEADERS,
            "通道名称": NORMALIZED_CHANNEL_NAME_HEADERS,
            "主码号": NORMALIZED_MAIN_CODE_HEADERS,
        },
        "通道编号映射表",
    )
    mapping: dict[tuple[str, str, str], list[str]] = {}
    for _row_number, row in rows[header_index + 1 :]:
        main_code = row[columns["主码号"]] if columns["主码号"] < len(row) else ""
        channel = row[columns["通道编号"]] if columns["通道编号"] < len(row) else ""
        channel_name = row[columns["通道名称"]] if columns["通道名称"] < len(row) else ""
        if not re.fullmatch(r"\d{8}", main_code) or not channel:
            continue
        for operator in ("移动", "联通", "电信"):
            if operator in channel_name:
                variant = classify_mobile_channel(channel_name) if operator == "移动" else ""
                values = mapping.setdefault((main_code, operator, variant), [])
                if channel not in values:
                    values.append(channel)
    return mapping


def load_channel_mapping_records(
    path: Path,
) -> dict[tuple[str, str], list[ChannelMappingRecord]]:
    """读取通道、运营商和省份；省份保留映射表原文，不在代码中写死。"""
    _sheet_name, rows = read_first_xlsx_sheet(path)
    header_index, columns = find_required_header(
        rows,
        {
            "通道编号": NORMALIZED_CHANNEL_CODE_HEADERS,
            "通道名称": NORMALIZED_CHANNEL_NAME_HEADERS,
            "主码号": NORMALIZED_MAIN_CODE_HEADERS,
            "省份": NORMALIZED_CHANNEL_PROVINCE_HEADERS,
        },
        "通道编号映射表",
    )
    mapping: dict[tuple[str, str], list[ChannelMappingRecord]] = {}
    for _row_number, row in rows[header_index + 1 :]:
        main_code = normalize_full_port(
            row[columns["主码号"]] if columns["主码号"] < len(row) else ""
        )
        channel = value_to_text(
            row[columns["通道编号"]] if columns["通道编号"] < len(row) else ""
        )
        channel_name = value_to_text(
            row[columns["通道名称"]] if columns["通道名称"] < len(row) else ""
        )
        province = value_to_text(
            row[columns["省份"]] if columns["省份"] < len(row) else ""
        )
        if not re.fullmatch(r"\d{8}", main_code) or not channel:
            continue
        for operator in ("移动", "联通", "电信"):
            if operator not in channel_name:
                continue
            record = ChannelMappingRecord(channel, province)
            values = mapping.setdefault((main_code, operator), [])
            if record not in values:
                values.append(record)
    return mapping


def load_signature_total(
    path: Path,
    expected_main_code: str,
) -> dict[str, list[SignatureTotalRecord]]:
    """读取签名总表表1，只索引文件名主码号下的子端口和业务类型。"""
    _sheet_name, rows = read_first_xlsx_sheet(path)
    header_index, columns = find_required_header(
        rows,
        {
            "完整子端口": NORMALIZED_SIGNATURE_TOTAL_PORT_HEADERS,
            "业务类型": NORMALIZED_SIGNATURE_TOTAL_BUSINESS_HEADERS,
        },
        f"签名总表“{path.name}”",
    )
    records: dict[str, list[SignatureTotalRecord]] = {}
    for row_number, row in rows[header_index + 1 :]:
        full_port = normalize_full_port(
            row[columns["完整子端口"]] if columns["完整子端口"] < len(row) else ""
        )
        if not full_port.startswith(expected_main_code):
            continue
        business_source = row[columns["业务类型"]] if columns["业务类型"] < len(row) else ""
        records.setdefault(full_port, []).append(
            SignatureTotalRecord(row_number, full_port, business_source, raw_row_text(row))
        )
    return records


def discover_signature_totals(
    data_dir: Path,
    region: str,
    operator: str,
    selected_folder: Path | None = None,
) -> dict[tuple[str, str], Path]:
    """从规范文件名动态取得主码号及普通/5G类型。"""
    folder = selected_folder.resolve() if selected_folder else data_dir / f"{region}{operator}"
    if not folder.exists() or not folder.is_dir():
        raise ProcessingError(f"未找到地区运营商文件夹：{folder.name}")
    pattern = re.compile(
        rf"^(\d{{8}})(-5G)?{re.escape(region)}{re.escape(operator)}签名总表\.xlsx$",
        re.IGNORECASE,
    )
    totals: dict[tuple[str, str], Path] = {}
    for path in folder.iterdir():
        if not path.is_file() or path.name.startswith("~$"):
            continue
        match = pattern.fullmatch(path.name)
        if match:
            variant = MOBILE_VARIANT_5G if match.group(2) else MOBILE_VARIANT_REGULAR
            key = (match.group(1), variant)
            if key in totals:
                raise ProcessingError(
                    f"主码号 {key[0]} 存在多个{variant}移动签名总表："
                    f"{totals[key].name}、{path.name}"
                )
            totals[key] = path
    if not totals:
        example = f"12345678{region}{operator}签名总表.xlsx 或 12345678-5G{region}{operator}签名总表.xlsx"
        raise ProcessingError(
            f"文件夹“{folder.name}”中没有符合命名规则的签名总表，例如：{example}"
        )
    return totals


def read_self_owned_records(
    source_path: Path,
    operator: str,
    important_support: str,
    data_dir: Path,
    signature_total_dir: Path | None = None,
) -> tuple[
    list[SelfOwnedOutputRecord],
    list[SelfOwnedErrorRecord],
    int,
    int,
    int,
]:
    """读取运营商已生效引流，只导入规范签名总表文件名声明的主码号。"""
    source_sheet, rows = read_first_xlsx_sheet(source_path)
    header_index, columns = find_required_header(
        rows,
        {
            "完整子端口": NORMALIZED_SELF_SOURCE_PORT_HEADERS,
            "归属省": NORMALIZED_SELF_SOURCE_PROVINCE_HEADERS,
            "审核状态": NORMALIZED_SELF_SOURCE_STATUS_HEADERS,
            "消息类型": NORMALIZED_SELF_SOURCE_MESSAGE_TYPE_HEADERS,
            "签名": NORMALIZED_SIGNATURE_HEADERS,
            "报备短链": NORMALIZED_SELF_SOURCE_SHORTLINK_HEADERS,
            "报备号码1": NORMALIZED_SELF_SOURCE_NUMBER1_HEADERS,
            "报备号码2": NORMALIZED_SELF_SOURCE_NUMBER2_HEADERS,
        },
        "运营商报备表",
    )

    mapping_path = channel_mapping_path(data_dir)
    if not mapping_path.exists():
        raise ProcessingError(
            f"未找到主码号资料“{CHANNEL_MAPPING_FILE}”。请重新安装或恢复本机资料。"
        )
    channel_mapping = load_channel_mapping(mapping_path)
    errors: list[SelfOwnedErrorRecord] = []
    duplicate_count = 0
    skipped_count = 0
    dirty_count = 0
    active_rows: list[tuple[int, list[str]]] = []
    active_regions: set[str] = set()
    for row_number, row in rows[header_index + 1 :]:
        if not any(row):
            continue
        status = row[columns["审核状态"]] if columns["审核状态"] < len(row) else ""
        if status != "已生效":
            skipped_count += 1
            continue
        active_rows.append((row_number, row))
        province = row[columns["归属省"]] if columns["归属省"] < len(row) else ""
        region = normalize_region_name(province)
        if region:
            active_regions.add(region)

    if not active_rows:
        raise ProcessingError("关键信息报备表中没有审核状态为“已生效”的数据。")
    if len(active_regions) != 1:
        displayed = "、".join(sorted(active_regions)) if active_regions else "全部为空"
        raise ProcessingError(f"已生效数据必须且只能包含一个归属省，当前为：{displayed}")
    region = next(iter(active_regions))
    signature_paths = discover_signature_totals(data_dir, region, operator, signature_total_dir)

    configuration_errors: list[str] = []
    channels_by_key: dict[tuple[str, str], str] = {}
    signature_indexes: dict[tuple[str, str], dict[str, list[SignatureTotalRecord]]] = {}
    for key, signature_path in sorted(signature_paths.items()):
        main_code, variant = key
        channels = channel_mapping.get((main_code, operator, variant), [])
        if not channels:
            configuration_errors.append(
                f"主码号 {main_code} 在通道映射表中没有{variant}{operator}通道"
            )
        elif len(channels) > 1:
            configuration_errors.append(
                f"主码号 {main_code} 在通道映射表中对应多个{variant}{operator}通道"
            )
        else:
            channels_by_key[key] = channels[0]
        signature_index = load_signature_total(signature_path, main_code)
        signature_indexes[key] = signature_index
        for full_port, total_records in signature_index.items():
            if len(total_records) <= 1:
                continue
            first_record = total_records[0]
            mapped_types = [map_business_type(item.business_source) for item in total_records]
            duplicate_kind = "业务类型一致，已合并使用" if len(set(mapped_types)) == 1 else "业务类型存在冲突"
            for duplicate in total_records[1:]:
                errors.append(
                    SelfOwnedErrorRecord(
                        signature_path.name,
                        duplicate.row_number,
                        full_port,
                        region,
                        duplicate.raw_content,
                        f"签名总表中完整子端口重复且{duplicate_kind}；首条位于第{first_record.row_number}行",
                    )
                )
    if configuration_errors:
        raise ProcessingError("配置检查失败：" + "；".join(configuration_errors))

    owned_main_codes = {main_code for main_code, _variant in signature_paths}
    grouped_rows: dict[
        tuple[str, str],
        list[tuple[int, str, str, tuple[str, ...], tuple[str, ...], str]],
    ] = {}
    for row_number, row in active_rows:
        full_port = row[columns["完整子端口"]] if columns["完整子端口"] < len(row) else ""
        province = row[columns["归属省"]] if columns["归属省"] < len(row) else ""
        raw = raw_row_text(row)
        full_port = normalize_full_port(full_port)
        row_region = normalize_region_name(province)
        if not full_port:
            errors.append(
                SelfOwnedErrorRecord(source_sheet, row_number, "", row_region, raw, "缺少完整子端口号")
            )
            continue
        if not re.fullmatch(r"\d{9,}", full_port):
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    "完整子端口号必须是纯数字，且长度必须大于 8 位",
                )
            )
            continue
        if not row_region:
            errors.append(
                SelfOwnedErrorRecord(source_sheet, row_number, full_port, "", raw, "缺少归属省")
            )
            continue
        main_code = full_port[:8]
        if main_code not in owned_main_codes:
            dirty_count += 1
            continue
        message_type_column = columns["消息类型"]
        raw_message_type = row[message_type_column] if message_type_column < len(row) else ""
        variant = map_mobile_message_type(raw_message_type)
        if not variant:
            displayed = value_to_text(raw_message_type) or "空白"
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    f"消息类型“{displayed}”无法识别，只支持“行短”或“5G消息”",
                )
            )
            continue
        route_key = (main_code, variant)
        if route_key not in signature_paths:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    f"未找到主码号 {main_code} 对应的{variant}移动签名总表",
                )
            )
            continue
        raw_signature = row[columns["签名"]] if columns["签名"] < len(row) else ""
        signature = normalize_signature(raw_signature)
        raw_link = row[columns["报备短链"]] if columns["报备短链"] < len(row) else ""
        raw_number1 = row[columns["报备号码1"]] if columns["报备号码1"] < len(row) else ""
        raw_number2 = row[columns["报备号码2"]] if columns["报备号码2"] < len(row) else ""
        cleaned_links: list[str] = []
        for item in split_items(raw_link):
            # 自家移动报备短链末尾的斜杠不写入生效短链结果。
            cleaned = item.rstrip("/")
            if cleaned and not is_blank_diversion(cleaned):
                cleaned_links.append(cleaned)
        links = tuple(ordered_unique(cleaned_links))
        numbers = tuple(
            ordered_unique(
                item
                for value in (raw_number1, raw_number2)
                for item in split_items(value)
                if not is_blank_diversion(item)
            )
        )
        grouped_rows.setdefault((full_port, variant), []).append(
            (row_number, signature, raw, links, numbers, row_region)
        )

    shortlink_records: list[SelfOwnedOutputRecord] = []
    for (full_port, variant), candidates in grouped_rows.items():
        fingerprints = {(item[1], item[3], item[4]) for item in candidates}
        if len(fingerprints) > 1:
            first = candidates[0]
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    first[0],
                    full_port,
                    first[5],
                    first[2],
                    "同一完整子端口在关键信息表中存在互相冲突的数据",
                )
            )
            continue
        duplicate_count += len(candidates) - 1
        row_number, signature, raw, links, numbers, row_region = candidates[0]
        if not signature:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet, row_number, full_port, row_region, raw, "关键信息表中的签名为空"
                )
            )
            continue
        if not links and not numbers:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    "报备短链、报备号码1、报备号码2均为空",
                )
            )
            continue
        invalid_numbers = [number for number in numbers if not is_number_item(number)]
        if invalid_numbers:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    "报备号码格式无效：" + "、".join(invalid_numbers),
                )
            )
            continue
        if len(links) > 1 or len(numbers) > 2:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    "同一完整子端口超过 1 个链接或 2 个号码，无法保持唯一子端口",
                )
            )
            continue

        main_code = full_port[:8]
        route_key = (main_code, variant)
        signature_matches = signature_indexes[route_key].get(full_port, [])
        if not signature_matches:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    f"{variant}移动签名总表表1中未找到完整子端口",
                )
            )
            continue
        mapped_types = [map_business_type(item.business_source) for item in signature_matches]
        if any(not value for value in mapped_types):
            unknown = ordered_unique(
                item.business_source or "空白"
                for item, mapped in zip(signature_matches, mapped_types)
                if not mapped
            )
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    "签名总表中的业务类型无法转换：" + "、".join(unknown),
                )
            )
            continue
        unique_types = ordered_unique(mapped_types)
        if len(unique_types) > 1:
            errors.append(
                SelfOwnedErrorRecord(
                    source_sheet,
                    row_number,
                    full_port,
                    row_region,
                    raw,
                    "签名总表中同一完整子端口对应多个业务类型",
                )
            )
            continue
        business_type = unique_types[0]
        sub_port = full_port[8:]
        shortlink_records.append(
            SelfOwnedOutputRecord(
                channels_by_key[route_key],
                signature,
                sub_port,
                operator,
                business_type,
                important_support,
                links[0] if links else "",
                ",".join(numbers),
            )
        )

    return shortlink_records, errors, duplicate_count, skipped_count, dirty_count


def read_first_csv_sheet(path: Path) -> tuple[str, list[tuple[int, list[str]]]]:
    """按既有编码兼容规则读取电信失败回执。"""
    text = decode_text_file(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = [
        (row_number, [value.strip() for value in row])
        for row_number, row in enumerate(csv.reader(io.StringIO(text, newline=""), dialect), 1)
    ]
    return path.name, rows


def read_first_tabular_sheet(path: Path) -> tuple[str, list[tuple[int, list[str]]]]:
    if path.suffix.lower() == ".csv":
        return read_first_csv_sheet(path)
    if path.suffix.lower() == ".xlsx":
        return read_first_xlsx_sheet(path)
    raise ProcessingError(f"电信失败结果不支持该格式：{path.suffix or '无扩展名'}")


def classify_telecom_xlsx(path: Path) -> str | None:
    """只查看首表前20行识别文件类型，避免为识别大平台表而整表加载。"""
    submission_required = {
        normalize_header("所属主端口"),
        normalize_header("子端口号"),
        normalize_header("业务类型"),
        normalize_header("子端口签名"),
    }
    platform_required = {
        normalize_header("主端口号"),
        normalize_header("子端口号"),
        normalize_header("企业签名"),
        normalize_header("加白链接"),
        normalize_header("加白号码"),
        normalize_header("状态"),
    }
    receipt_main_headers = (
        NORMALIZED_TELECOM_RECEIPT_MAIN_HEADERS
        | NORMALIZED_TELECOM_SUBMISSION_MAIN_HEADERS
    )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        reset_input_workbook_dimensions(workbook)
        sheet = workbook.worksheets[0]
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_number > 20:
                break
            headers = {normalize_header(value) for value in cells if value_to_text(value)}
            if platform_required.issubset(headers):
                return "platform"
            if (
                headers & NORMALIZED_TELECOM_RECEIPT_REASON_HEADERS
                and headers & NORMALIZED_TELECOM_RECEIPT_PORT_HEADERS
                and headers & receipt_main_headers
            ):
                return "receipt"
            if submission_required.issubset(headers):
                return "submission"
        return None
    finally:
        workbook.close()


def telecom_submission_filename_role(path: Path) -> str:
    """只从文件名读取“引流/签名”角色，不用文件名决定所属省份。"""
    compact = re.sub(r"\s+", "", path.stem)
    has_diversion = "引流" in compact
    has_signature = "签名" in compact
    if has_diversion == has_signature:
        return TELECOM_CATEGORY_MIXED
    return TELECOM_CATEGORY_DIVERSION if has_diversion else TELECOM_CATEGORY_SIGNATURE


def discover_telecom_batch(
    batch_dir: Path,
) -> tuple[list[Path], list[Path], Path]:
    if not batch_dir.exists() or not batch_dir.is_dir():
        raise ProcessingError("请选择存在的自家电信批次文件夹。")
    xlsx_files = sorted(
        path for path in batch_dir.iterdir()
        if path.is_file() and path.suffix.lower() == ".xlsx" and not path.name.startswith("~$")
    )
    receipts: list[Path] = sorted(
        path for path in batch_dir.iterdir()
        if path.is_file() and path.suffix.lower() == ".csv" and not path.name.startswith("~$")
    )
    submissions: list[Path] = []
    platforms: list[Path] = []
    for path in xlsx_files:
        kind = classify_telecom_xlsx(path)
        if kind == "submission":
            submissions.append(path)
        elif kind == "platform":
            platforms.append(path)
        elif kind == "receipt":
            receipts.append(path)
    if not submissions:
        raise ProcessingError("批次文件夹中没有识别到电信提交报备表。")
    if len(platforms) != 1:
        raise ProcessingError(
            f"批次文件夹中必须且只能有 1 份平台引流子码表，当前识别到 {len(platforms)} 份。"
        )
    if not receipts:
        raise ProcessingError(
            "批次文件夹中没有识别到运营商失败结果。即使没有失败数据，也必须放入只有表头的回执。"
        )

    return submissions, sorted(receipts), platforms[0]


def read_telecom_submission(
    path: Path,
    category: str = TELECOM_CATEGORY_MIXED,
) -> list[TelecomSubmissionRecord]:
    sheet_name, rows = read_first_xlsx_sheet(path)
    header_index, columns = find_required_header(
        rows,
        {
            "主码号": NORMALIZED_TELECOM_SUBMISSION_MAIN_HEADERS,
            "完整子端口": NORMALIZED_TELECOM_SUBMISSION_PORT_HEADERS,
            "业务类型": NORMALIZED_TELECOM_SUBMISSION_BUSINESS_HEADERS,
            "签名": NORMALIZED_TELECOM_SUBMISSION_SIGNATURE_HEADERS,
        },
        f"电信提交表“{path.name}”",
    )
    header_row = rows[header_index][1]
    enterprise_column = find_optional_column(
        header_row, NORMALIZED_TELECOM_SUBMISSION_ENTERPRISE_HEADERS
    )
    message_column = find_optional_column(
        header_row, NORMALIZED_TELECOM_SUBMISSION_MESSAGE_HEADERS
    )
    records: list[TelecomSubmissionRecord] = []
    for row_number, row in rows[header_index + 1 :]:
        main_code = normalize_full_port(row[columns["主码号"]] if columns["主码号"] < len(row) else "")
        full_port = normalize_full_port(
            row[columns["完整子端口"]] if columns["完整子端口"] < len(row) else ""
        )
        if not re.fullmatch(r"\d{8}", main_code) or not re.fullmatch(r"\d{9,}", full_port):
            continue
        signature = normalize_signature(row[columns["签名"]] if columns["签名"] < len(row) else "")
        business_type = map_business_type(
            row[columns["业务类型"]] if columns["业务类型"] < len(row) else ""
        )
        enterprise_name = value_to_text(
            row[enterprise_column]
            if enterprise_column is not None and enterprise_column < len(row)
            else ""
        )
        message_template = value_to_text(
            row[message_column]
            if message_column is not None and message_column < len(row)
            else ""
        )
        records.append(
            TelecomSubmissionRecord(
                f"{path.name} / {sheet_name}",
                row_number,
                main_code,
                full_port,
                signature,
                business_type,
                enterprise_name,
                message_template,
                category,
                raw_row_text(row),
            )
        )
    if not records:
        raise ProcessingError(f"电信提交表“{path.name}”中没有有效业务数据。")
    return records


def read_and_categorize_telecom_submissions(
    paths: Sequence[Path],
    channel_mapping: dict[tuple[str, str], list[ChannelMappingRecord]],
) -> list[TelecomSubmissionRecord]:
    """用主码号映射省份确认湖南，再用文件名区分湖南引流/签名角色。"""
    all_records: list[TelecomSubmissionRecord] = []
    file_categories: set[str] = set()
    for path in paths:
        records = read_telecom_submission(path)
        provinces: set[str] = set()
        for main_code in {
            record.main_code
            for record in records
        }:
            matches = channel_mapping.get((main_code, "电信"), [])
            if len(matches) == 1 and matches[0].province:
                provinces.add(normalize_region_name(matches[0].province))
        if len(provinces) > 1:
            raise ProcessingError(
                f"提交表“{path.name}”中的主码号在通道映射表中对应多个省份："
                + "、".join(sorted(provinces))
            )

        filename_role = telecom_submission_filename_role(path)
        is_hunan = provinces == {"湖南"}
        if "湖南" in path.stem and provinces and not is_hunan:
            raise ProcessingError(
                f"提交表“{path.name}”文件名写有湖南，但主码号映射省份不是湖南。"
            )
        if is_hunan:
            if filename_role == TELECOM_CATEGORY_MIXED:
                raise ProcessingError(
                    f"湖南提交表“{path.name}”的文件名必须明确包含“引流”或“签名”，且只能包含一个。"
                )
            category = filename_role
        else:
            category = TELECOM_CATEGORY_MIXED
        file_categories.add(category)
        all_records.extend(replace(record, category=category) for record in records)

    if TELECOM_CATEGORY_MIXED in file_categories and len(file_categories) > 1:
        raise ProcessingError(
            "湖南引流/签名分表不能与其他地区的普通电信混合报备表放在同一批次文件夹。"
        )
    return all_records


def read_telecom_failed_records(
    path: Path,
) -> tuple[list[TelecomReceiptFailureRecord], set[str]]:
    source, rows = read_first_tabular_sheet(path)
    header_index, columns = find_required_header(
        rows,
        {
            "完整子端口": NORMALIZED_TELECOM_RECEIPT_PORT_HEADERS,
        },
        f"电信失败回执“{path.name}”",
    )
    header_row = rows[header_index][1]
    main_column = find_optional_column(
        header_row,
        NORMALIZED_TELECOM_RECEIPT_MAIN_HEADERS
        | NORMALIZED_TELECOM_SUBMISSION_MAIN_HEADERS,
    )
    reason_column = find_optional_column(
        header_row,
        NORMALIZED_TELECOM_RECEIPT_REASON_HEADERS,
    )
    records: list[TelecomReceiptFailureRecord] = []
    main_codes: set[str] = set()
    for row_number, row in rows[header_index + 1 :]:
        full_port = normalize_full_port(
            row[columns["完整子端口"]] if columns["完整子端口"] < len(row) else ""
        )
        if not re.fullmatch(r"\d{9,}", full_port):
            continue
        main_code = normalize_full_port(
            row[main_column]
            if main_column is not None and main_column < len(row)
            else full_port[:8]
        )
        if main_code:
            main_codes.add(main_code)
        reason = value_to_text(
            row[reason_column]
            if reason_column is not None and reason_column < len(row)
            else ""
        )
        records.append(
            TelecomReceiptFailureRecord(
                source=source,
                row_number=row_number,
                main_code=main_code or full_port[:8],
                full_port=full_port,
                reason=reason,
            )
        )
    return records, main_codes


def read_telecom_failed_ports(path: Path) -> tuple[set[str], set[str]]:
    """兼容原有调用：返回失败完整子端口集合及主码号集合。"""
    records, main_codes = read_telecom_failed_records(path)
    return {record.full_port for record in records}, main_codes


def read_telecom_platform(path: Path) -> dict[str, list[TelecomPlatformRecord]]:
    _sheet_name, rows = read_first_xlsx_sheet(path)
    header_index, columns = find_required_header(
        rows,
        {
            "完整子端口": NORMALIZED_TELECOM_PLATFORM_PORT_HEADERS,
            "签名": NORMALIZED_TELECOM_PLATFORM_SIGNATURE_HEADERS,
            "加白链接": NORMALIZED_TELECOM_PLATFORM_LINK_HEADERS,
            "加白号码": NORMALIZED_TELECOM_PLATFORM_NUMBER_HEADERS,
        },
        f"平台引流子码表“{path.name}”",
    )
    by_suffix: dict[str, list[TelecomPlatformRecord]] = {}
    for row_number, row in rows[header_index + 1 :]:
        full_port = normalize_full_port(
            row[columns["完整子端口"]] if columns["完整子端口"] < len(row) else ""
        )
        if not re.fullmatch(r"\d{9,}", full_port):
            continue
        suffix = full_port[8:]
        raw_link = row[columns["加白链接"]] if columns["加白链接"] < len(row) else ""
        raw_number = row[columns["加白号码"]] if columns["加白号码"] < len(row) else ""
        by_suffix.setdefault(suffix, []).append(
            TelecomPlatformRecord(row_number, suffix, raw_link, raw_number)
        )
    return by_suffix


def read_self_telecom_records(
    batch_dir: Path,
    important_support: str,
    data_dir: Path,
) -> tuple[
    list[SelfOwnedOutputRecord],
    list[SelfOwnedOutputRecord],
    list[FinalDiversionRecord],
    list[TelecomErrorRecord],
    list[TelecomFailureReasonRecord],
    int,
    int,
    int,
    int,
    int,
    list[Path],
]:
    submissions, receipt_paths, platform_path = discover_telecom_batch(batch_dir)
    mapping_path = channel_mapping_path(data_dir)
    if not mapping_path.exists():
        raise ProcessingError(f"未找到“{CHANNEL_MAPPING_FILE}”。")
    channel_mapping = load_channel_mapping_records(mapping_path)
    platform = read_telecom_platform(platform_path)

    errors: list[TelecomErrorRecord] = []
    submitted_records = read_and_categorize_telecom_submissions(
        submissions,
        channel_mapping,
    )
    submitted_by_full_port: dict[str, list[TelecomSubmissionRecord]] = {}
    for record in submitted_records:
        submitted_by_full_port.setdefault(record.full_port, []).append(record)
    submitted_ports = {record.full_port for record in submitted_records}
    failed_ports: set[str] = set()
    receipt_failures: list[TelecomReceiptFailureRecord] = []
    for receipt_path in receipt_paths:
        receipt_records, _receipt_mains = read_telecom_failed_records(receipt_path)
        receipt_failures.extend(receipt_records)
        failed_ports.update(record.full_port for record in receipt_records)
    matched_failed_ports = failed_ports & submitted_ports
    unrelated_failed_count = len(failed_ports - submitted_ports)
    failed_count = len(matched_failed_ports)
    failure_reason_records: list[TelecomFailureReasonRecord] = []
    reasons_by_port: dict[str, list[str]] = {}
    for failure in receipt_failures:
        reasons = reasons_by_port.setdefault(failure.full_port, [])
        if failure.reason not in reasons:
            reasons.append(failure.reason)
    # 按提交顺序汇总运营商状态；完整子端口去重，不跨主码号合并。
    # 上报成功与本地导入校验是两个概念，不能用最终输出记录反推状态。
    for full_port, matches in submitted_by_full_port.items():
        submission = next(
            (item for item in matches if item.category == TELECOM_CATEGORY_DIVERSION),
            matches[0],
        )
        failed = full_port in matched_failed_ports
        for reason in reasons_by_port.get(full_port, [""]):
            failure_reason_records.append(TelecomFailureReasonRecord(
                main_code=submission.main_code,
                full_port=full_port,
                signature=signature_without_brackets(submission.signature),
                reason=reason,
                status="失败" if failed else "成功",
            ))
    successful_candidates = [
        record for record in submitted_records if record.full_port not in matched_failed_ports
    ]

    # 与自家移动保持一致：映射表中没有本运营商通道的主码号属于脏数据，
    # 直接忽略，不让单个外部主码号导致整个电信批次停止。
    dirty_count = 0
    mapped_candidates: list[TelecomSubmissionRecord] = []
    for record in successful_candidates:
        if channel_mapping.get((record.main_code, "电信"), []):
            mapped_candidates.append(record)
        else:
            dirty_count += 1
    successful_candidates = mapped_candidates

    grouped: dict[str, list[TelecomSubmissionRecord]] = {}
    for record in successful_candidates:
        grouped.setdefault(record.full_port, []).append(record)
    duplicate_count = 0
    unique_success: list[TelecomSubmissionRecord] = []
    for full_port, candidates in grouped.items():
        diversion_candidates = [
            item for item in candidates if item.category == TELECOM_CATEGORY_DIVERSION
        ]
        if diversion_candidates:
            chosen_candidates = diversion_candidates
        else:
            chosen_candidates = candidates
        # 失败回执已经先按完整子端口剔除，因此这里剩余的重复记录均视为成功。
        # 无论签名、业务类型、企业名称或短信模板是否不同，都按稳定的提交顺序
        # 保留一条；湖南同一完整子端口同时出现在引流/签名表时仍优先保留引流。
        retained_candidate = chosen_candidates[0]
        duplicate_count += len(candidates) - 1
        for duplicate in candidates:
            if duplicate is retained_candidate:
                continue
            errors.append(
                TelecomErrorRecord(
                    duplicate.source,
                    duplicate.row_number,
                    duplicate.main_code,
                    duplicate.full_port,
                    duplicate.signature,
                    duplicate.raw_content,
                    "同一完整子端口重复，已按规则去重保留首条",
                )
            )
        unique_success.append(retained_candidate)

    main_codes = {record.main_code for record in unique_success}
    channels_by_main: dict[str, ChannelMappingRecord] = {}
    configuration_errors: list[str] = []
    for main_code in sorted(main_codes):
        channel_records = channel_mapping.get((main_code, "电信"), [])
        if not channel_records:
            configuration_errors.append(f"主码号 {main_code} 没有电信通道")
        elif len(channel_records) > 1:
            configuration_errors.append(f"主码号 {main_code} 对应多个电信通道")
        else:
            channels_by_main[main_code] = channel_records[0]
    if configuration_errors:
        raise ProcessingError("配置检查失败：" + "；".join(configuration_errors))

    shortlinks: list[SelfOwnedOutputRecord] = []
    signatures: list[SelfOwnedOutputRecord] = []
    final_diversions: list[FinalDiversionRecord] = []
    for record in unique_success:
        if not record.full_port.startswith(record.main_code):
            errors.append(
                TelecomErrorRecord(
                    record.source,
                    record.row_number,
                    record.main_code,
                    record.full_port,
                    record.signature,
                    record.raw_content,
                    "完整子端口号的前8位与所属主端口不一致",
                )
            )
            continue
        if not record.signature:
            errors.append(
                TelecomErrorRecord(
                    record.source, record.row_number, record.main_code, record.full_port,
                    "", record.raw_content, "提交表中的子端口签名为空"
                )
            )
            continue
        if not record.business_type:
            errors.append(
                TelecomErrorRecord(
                    record.source, record.row_number, record.main_code, record.full_port,
                    record.signature, record.raw_content, "业务类型不是通知提醒或广告促销"
                )
            )
            continue
        suffix = record.full_port[8:]
        platform_matches = platform.get(suffix, [])
        is_diversion = (
            record.category == TELECOM_CATEGORY_DIVERSION
            or (record.category == TELECOM_CATEGORY_MIXED and bool(platform_matches))
        )
        output_base = {
            "channel": channels_by_main[record.main_code].channel,
            "signature": record.signature,
            "sub_port": suffix,
            "operator": "电信",
            "business_type": record.business_type,
            "important_support": important_support,
        }
        if not is_diversion:
            signatures.append(SelfOwnedOutputRecord(**output_base))
            continue
        if not platform_matches:
            errors.append(
                TelecomErrorRecord(
                    record.source,
                    record.row_number,
                    record.main_code,
                    record.full_port,
                    record.signature,
                    record.raw_content,
                    "湖南引流报备成功，但平台引流子码表中未找到对应子端口后缀",
                )
            )
            continue
        links = ordered_unique(
            item
            for match in platform_matches
            for item in split_items(match.raw_link)
            if not is_blank_diversion(item)
        )
        numbers = ordered_unique(
            item
            for match in platform_matches
            for item in split_items(match.raw_number)
            if not is_blank_diversion(item)
        )
        invalid_numbers = [number for number in numbers if not is_number_item(number)]
        if invalid_numbers:
            errors.append(
                TelecomErrorRecord(
                    record.source, record.row_number, record.main_code, record.full_port,
                    record.signature, record.raw_content,
                    "平台加白号码格式无效：" + "、".join(invalid_numbers)
                )
            )
            continue
        if not links and not numbers:
            errors.append(
                TelecomErrorRecord(
                    record.source, record.row_number, record.main_code, record.full_port,
                    record.signature, record.raw_content,
                    "已匹配平台引流子码，但加白链接和加白号码都为空"
                )
            )
            continue
        if len(links) > 1 or len(numbers) > 2:
            errors.append(
                TelecomErrorRecord(
                    record.source, record.row_number, record.main_code, record.full_port,
                    record.signature, record.raw_content,
                    "平台同一子端口超过1个链接或2个号码，无法保持唯一子端口"
                )
            )
            continue
        final_validation_errors: list[str] = []
        province = channels_by_main[record.main_code].province
        if not province:
            final_validation_errors.append("通道映射表中该电信主码号的省份为空")
        if not record.enterprise_name:
            final_validation_errors.append("提交表中的企业名称为空")
        if not record.message_template:
            final_validation_errors.append("提交表中的短信模板为空")
        else:
            message_signature = re.match(r"^【([^】]+)】", record.message_template)
            plain_signature = signature_without_brackets(record.signature)
            if message_signature is None:
                final_validation_errors.append("短信模板没有以【签名】开头")
            elif message_signature.group(1).strip() != plain_signature:
                final_validation_errors.append("短信模板开头签名与子端口签名不一致")
        if final_validation_errors:
            errors.append(
                TelecomErrorRecord(
                    record.source,
                    record.row_number,
                    record.main_code,
                    record.full_port,
                    record.signature,
                    record.raw_content,
                    "；".join(final_validation_errors),
                )
            )
            continue
        shortlinks.append(
            SelfOwnedOutputRecord(
                **output_base,
                link=links[0] if links else "",
                number=",".join(numbers),
            )
        )
        final_diversions.append(
            FinalDiversionRecord(
                province=province,
                main_code=record.main_code,
                full_port=record.full_port,
                enterprise_name=record.enterprise_name,
                number1=numbers[0] if numbers else "",
                number2=numbers[1] if len(numbers) > 1 else "",
                link=links[0] if links else "",
                signature=signature_without_brackets(record.signature),
                content=record.message_template,
            )
        )
    input_paths = [
        platform_path,
        mapping_path,
        *submissions,
        *receipt_paths,
    ]
    return (
        shortlinks,
        signatures,
        final_diversions,
        errors,
        failure_reason_records,
        duplicate_count,
        failed_count,
        len(submissions),
        unrelated_failed_count,
        dirty_count,
        input_paths,
    )


def find_template_header(worksheet: object, required: Sequence[str]) -> tuple[int, dict[str, int]]:
    required_set = set(required)
    for row_number in range(1, min(getattr(worksheet, "max_row", 1), 10) + 1):
        mapping: dict[str, int] = {}
        for cell in worksheet[row_number]:
            if cell.value is not None:
                mapping[value_to_text(cell.value)] = cell.column
        if required_set.issubset(mapping):
            return row_number, mapping
    missing = "、".join(required)
    raise ProcessingError(f"模板表头不正确。应包含：{missing}")


def ensure_named_style(workbook: object, name: str, number_format: str) -> str:
    if name in getattr(workbook, "named_styles", []):
        return name
    style = NamedStyle(name=name)
    style.font = Font(name="宋体", size=11)
    style.alignment = Alignment(vertical="center")
    style.number_format = number_format
    workbook.add_named_style(style)
    return name


def snapshot_row_styles(worksheet: object, row_number: int, max_column: int) -> list[dict[str, object]]:
    snapshots: list[dict[str, object]] = []
    for column in range(1, max_column + 1):
        cell = worksheet.cell(row_number, column)
        snapshots.append(
            {
                "style": copy(cell._style),
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
                "number_format": cell.number_format,
            }
        )
    return snapshots


def apply_snapshot(cell: Cell, snapshot: dict[str, object]) -> None:
    cell._style = copy(snapshot["style"])
    cell.font = copy(snapshot["font"])
    cell.fill = copy(snapshot["fill"])
    cell.border = copy(snapshot["border"])
    cell.alignment = copy(snapshot["alignment"])
    cell.protection = copy(snapshot["protection"])
    cell.number_format = str(snapshot["number_format"])


def write_template_output(
    template_path: Path,
    output_path: Path,
    mode: str,
    records: Sequence[tuple[str, ...]],
    channel: str,
    start_sub_port: int,
    operator: str,
    business_type: str,
    important_support: str,
) -> None:
    """复制外置模板并写入正式数据，不修改原模板文件。"""
    try:
        workbook = load_workbook(template_path)
    except Exception as exc:
        raise ProcessingError(f"无法读取模板“{template_path.name}”：{exc}") from exc

    try:
        worksheet = workbook["导入模板"] if "导入模板" in workbook.sheetnames else workbook.active
        required = SHORTLINK_OUTPUT_HEADERS if mode == MODE_SHORTLINK else SIGNATURE_OUTPUT_HEADERS
        header_row, columns = find_template_header(worksheet, required)
        data_start = header_row + 1
        max_output_column = max(columns[header] for header in required)
        style_snapshots = None
        if worksheet.max_row >= data_start:
            style_snapshots = snapshot_row_styles(worksheet, data_start, max_output_column)
            worksheet.delete_rows(data_start, worksheet.max_row - data_start + 1)

        general_style = ensure_named_style(workbook, "导入数据", "General")
        text_style = ensure_named_style(workbook, "导入数据文本", "@")
        integer_style = ensure_named_style(workbook, "导入数据整数", "0")

        preferred_widths = (
            {
                "通道编号": 14,
                "签名": 22,
                "子端口": 14,
                "运营商": 10,
                "业务类型": 10,
                "等级": 8,
                "链接": 32,
                "号码": 28,
                "是否重保": 10,
                "发送量": 10,
                "错误信息": 18,
                "备注": 18,
            }
            if mode == MODE_SHORTLINK
            else {
                "通道编号": 14,
                "签名": 22,
                "子端口": 14,
                "运营商": 10,
                "业务类型": 10,
                "是否重保": 10,
                "备注": 18,
            }
        )
        for header, width in preferred_widths.items():
            letter = get_column_letter(columns[header])
            current_width = worksheet.column_dimensions[letter].width or 0
            worksheet.column_dimensions[letter].width = max(current_width, width)

        for offset, record in enumerate(records):
            row_number = data_start + offset
            signature = record[0]
            values: dict[str, object]
            if mode == MODE_SHORTLINK:
                link = record[1]
                number = record[2]
                values = {
                    "通道编号": channel,
                    "签名": signature,
                    "子端口": start_sub_port + offset,
                    "运营商": operator,
                    "业务类型": business_type,
                    "等级": 2,
                    "链接": link or None,
                    "号码": number or None,
                    "是否重保": important_support,
                }
            else:
                values = {
                    "通道编号": channel,
                    "签名": signature,
                    "子端口": start_sub_port + offset,
                    "运营商": operator,
                    "业务类型": business_type,
                    "是否重保": important_support,
                }

            for header, value in values.items():
                if value is None:
                    continue
                column = columns[header]
                cell = worksheet.cell(row_number, column, value)
                if style_snapshots is not None:
                    apply_snapshot(cell, style_snapshots[column - 1])
                elif header in {"通道编号", "签名", "链接", "号码"}:
                    cell.style = text_style
                elif header in {"子端口", "等级"}:
                    cell.style = integer_style
                else:
                    cell.style = general_style

        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    except PermissionError as exc:
        raise ProcessingError(
            f"无法保存结果文件。请关闭已打开的同名 Excel 文件后重试：{output_path.name}"
        ) from exc
    finally:
        workbook.close()


def write_self_owned_output(
    template_path: Path,
    output_path: Path,
    mode: str,
    records: Sequence[SelfOwnedOutputRecord],
) -> None:
    """写入已经逐行确定通道、子端口和业务类型的自家通道结果。"""
    try:
        workbook = load_workbook(template_path)
    except Exception as exc:
        raise ProcessingError(f"无法读取模板“{template_path.name}”：{exc}") from exc

    try:
        worksheet = workbook["导入模板"] if "导入模板" in workbook.sheetnames else workbook.active
        required = SHORTLINK_OUTPUT_HEADERS if mode == MODE_SHORTLINK else SIGNATURE_OUTPUT_HEADERS
        header_row, columns = find_template_header(worksheet, required)
        data_start = header_row + 1
        max_output_column = max(columns[header] for header in required)
        style_snapshots = None
        if worksheet.max_row >= data_start:
            style_snapshots = snapshot_row_styles(worksheet, data_start, max_output_column)
            worksheet.delete_rows(data_start, worksheet.max_row - data_start + 1)

        general_style = ensure_named_style(workbook, "导入数据", "General")
        text_style = ensure_named_style(workbook, "导入数据文本", "@")
        integer_style = ensure_named_style(workbook, "导入数据整数", "0")
        preferred_widths = (
            {
                "通道编号": 14,
                "签名": 22,
                "子端口": 14,
                "运营商": 10,
                "业务类型": 10,
                "等级": 8,
                "链接": 32,
                "号码": 28,
                "是否重保": 10,
                "发送量": 10,
                "错误信息": 18,
                "备注": 18,
            }
            if mode == MODE_SHORTLINK
            else {
                "通道编号": 14,
                "签名": 22,
                "子端口": 14,
                "运营商": 10,
                "业务类型": 10,
                "是否重保": 10,
                "备注": 18,
            }
        )
        for header, width in preferred_widths.items():
            letter = get_column_letter(columns[header])
            current_width = worksheet.column_dimensions[letter].width or 0
            worksheet.column_dimensions[letter].width = max(current_width, width)

        for offset, record in enumerate(records):
            row_number = data_start + offset
            values: dict[str, object] = {
                "通道编号": record.channel,
                "签名": record.signature,
                "子端口": record.sub_port,
                "运营商": record.operator,
                "业务类型": record.business_type,
                "是否重保": record.important_support,
            }
            if mode == MODE_SHORTLINK:
                values.update(
                    {
                        "等级": 2,
                        "链接": record.link or None,
                        "号码": record.number or None,
                    }
                )
            for header, value in values.items():
                if value is None:
                    continue
                column = columns[header]
                cell = worksheet.cell(row_number, column, value)
                if style_snapshots is not None:
                    apply_snapshot(cell, style_snapshots[column - 1])
                if header in {"通道编号", "签名", "子端口", "链接", "号码"}:
                    if style_snapshots is None:
                        cell.style = text_style
                    else:
                        cell.number_format = "@"
                elif header == "等级":
                    if style_snapshots is None:
                        cell.style = integer_style
                    else:
                        cell.number_format = "0"
                elif style_snapshots is None:
                    cell.style = general_style

        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    except PermissionError as exc:
        raise ProcessingError(
            f"无法保存结果文件。请关闭已打开的同名 Excel 文件后重试：{output_path.name}"
        ) from exc
    finally:
        workbook.close()


def find_final_diversion_template_header(
    workbook: object,
) -> tuple[object, int, dict[str, int]]:
    required = {
        header: normalize_header(header) for header in FINAL_DIVERSION_OUTPUT_HEADERS
    }
    for worksheet in workbook.worksheets:
        for row_number in range(1, min(getattr(worksheet, "max_row", 1), 10) + 1):
            normalized_columns = {
                normalize_header(cell.value): cell.column
                for cell in worksheet[row_number]
                if value_to_text(cell.value)
            }
            if all(normalized in normalized_columns for normalized in required.values()):
                return (
                    worksheet,
                    row_number,
                    {
                        header: normalized_columns[normalized]
                        for header, normalized in required.items()
                    },
                )
    raise ProcessingError(
        "最终报备表模板表头不正确，必须包含："
        + "、".join(FINAL_DIVERSION_OUTPUT_HEADERS)
        + "。"
    )


def write_final_diversion_output(
    template_path: Path,
    output_path: Path,
    records: Sequence[FinalDiversionRecord],
) -> None:
    """清除模板示例数据，写入与生效短链完全一致的成功引流记录。"""
    try:
        workbook = load_workbook(template_path)
    except Exception as exc:
        raise ProcessingError(f"无法读取模板“{template_path.name}”：{exc}") from exc

    try:
        worksheet, header_row, columns = find_final_diversion_template_header(workbook)
        data_start = header_row + 1
        max_output_column = max(columns.values())
        style_snapshots = None
        template_row_height = worksheet.row_dimensions[data_start].height
        if worksheet.max_row >= data_start:
            style_snapshots = snapshot_row_styles(worksheet, data_start, max_output_column)

        purpose = FINAL_DIVERSION_DEFAULT_PURPOSE
        operation: object = FINAL_DIVERSION_DEFAULT_OPERATION
        for row_number in range(data_start, worksheet.max_row + 1):
            purpose_value = value_to_text(
                worksheet.cell(row_number, columns["用途说明"]).value
            )
            operation_value = worksheet.cell(
                row_number, columns[FINAL_DIVERSION_OPERATION_HEADER]
            ).value
            if purpose_value:
                purpose = purpose_value
            if value_to_text(operation_value):
                operation = operation_value
            if purpose_value and value_to_text(operation_value):
                break

        if worksheet.max_row >= data_start:
            worksheet.delete_rows(data_start, worksheet.max_row - data_start + 1)

        text_headers = {
            "省份",
            "主端口号",
            "子端口号",
            "子端口企业名称",
            "号码1",
            "号码2",
            "链接",
            "用途说明",
            "签名",
            "内容",
        }
        for offset, record in enumerate(records):
            row_number = data_start + offset
            values: dict[str, object] = {
                "序号": offset + 1,
                "省份": record.province,
                "主端口号": record.main_code,
                "子端口号": record.full_port,
                "子端口企业名称": record.enterprise_name,
                "号码1": record.number1 or None,
                "号码2": record.number2 or None,
                "链接": record.link or None,
                "用途说明": purpose,
                FINAL_DIVERSION_OPERATION_HEADER: operation,
                "签名": record.signature,
                "内容": record.content,
            }
            if template_row_height is not None:
                worksheet.row_dimensions[row_number].height = template_row_height
            for header in FINAL_DIVERSION_OUTPUT_HEADERS:
                column = columns[header]
                cell = worksheet.cell(row_number, column, values.get(header))
                if style_snapshots is not None:
                    apply_snapshot(cell, style_snapshots[column - 1])
                else:
                    cell.font = Font(name="宋体", size=11)
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=header in {"子端口企业名称", "用途说明", "内容"},
                    )
                if header in text_headers:
                    cell.number_format = "@"
                    cell.quotePrefix = True
                elif header in {"序号", FINAL_DIVERSION_OPERATION_HEADER}:
                    cell.number_format = "0"

        if records:
            worksheet.auto_filter.ref = (
                f"{get_column_letter(min(columns.values()))}{header_row}:"
                f"{get_column_letter(max(columns.values()))}{data_start + len(records) - 1}"
            )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    except PermissionError as exc:
        raise ProcessingError(
            f"无法保存结果文件。请关闭已打开的同名 Excel 文件后重试：{output_path.name}"
        ) from exc
    finally:
        workbook.close()


def save_error_report(
    path: Path,
    errors: Sequence[ErrorRecord],
    include_category: bool = False,
) -> None:
    """只有存在异常时才创建错误明细；混合模式额外写入分流类型。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "错误明细"
    headers = (
        ["分流类型", "工作表/来源", "原始行号", "签名", "原始内容", "错误原因"]
        if include_category
        else ["工作表/来源", "原始行号", "签名", "原始内容", "错误原因"]
    )
    worksheet.append(headers)
    for error in errors:
        values = [error.source, error.row_number, error.signature, error.raw_content, error.reason]
        if include_category:
            values.insert(0, error.category)
        worksheet.append(values)

    header_fill = PatternFill("solid", fgColor="C00000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    last_column = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(1, worksheet.max_row)}"
    widths = [14, 18, 12, 28, 72, 32] if include_category else [18, 12, 28, 72, 32]
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()
    os.replace(temporary_path, path)


def save_self_owned_error_report(
    path: Path,
    errors: Sequence[SelfOwnedErrorRecord],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "错误明细"
    headers = ["工作表/来源", "原始行号", "完整子端口", "归属省", "原始内容", "错误原因"]
    worksheet.append(headers)
    for error in errors:
        worksheet.append(
            [
                error.source,
                error.row_number,
                error.full_port,
                error.province,
                error.raw_content,
                error.reason,
            ]
        )
    header_fill = PatternFill("solid", fgColor="C00000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:F{max(1, worksheet.max_row)}"
    for index, width in enumerate((18, 12, 24, 12, 72, 38), 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[2].number_format = "@"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()
    os.replace(temporary_path, path)


def save_telecom_error_report(
    path: Path,
    errors: Sequence[TelecomErrorRecord],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "错误明细"
    headers = ["工作表/来源", "原始行号", "主码号", "完整子端口", "签名", "原始内容", "错误原因"]
    worksheet.append(headers)
    for error in errors:
        worksheet.append(
            [
                error.source,
                error.row_number,
                error.main_code,
                error.full_port,
                error.signature,
                error.raw_content,
                error.reason,
            ]
        )
    header_fill = PatternFill("solid", fgColor="C00000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:G{max(1, worksheet.max_row)}"
    for index, width in enumerate((22, 12, 12, 24, 26, 72, 42), 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[2].number_format = "@"
        row[3].number_format = "@"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()
    os.replace(temporary_path, path)


def save_telecom_failure_reason_report(
    path: Path,
    records: Sequence[TelecomFailureReasonRecord],
) -> None:
    """输出全部提交子端口的运营商上报状态和原始错误原因。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "错误原因"
    headers = ["主码号", "子端口号", "签名", "上报状态", "错误原因"]
    worksheet.append(headers)
    for record in records:
        worksheet.append(
            [record.main_code, record.full_port, record.signature, record.status, record.reason]
        )

    header_fill = PatternFill("solid", fgColor="C00000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:E{max(1, worksheet.max_row)}"
    for index, width in enumerate((14, 24, 28, 14, 56), 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].number_format = "@"
        row[1].number_format = "@"

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()
    os.replace(temporary_path, path)


def choose_output_paths(source: Path, mode: str, include_error: bool) -> tuple[Path, Path | None]:
    """为单一模式选择不会覆盖旧文件的输出名称。"""
    label = "生效短链" if mode == MODE_SHORTLINK else "生效签名"
    index = 0
    while True:
        suffix = "" if index == 0 else f"({index})"
        result = source.with_name(f"{source.stem}_{label}导入结果{suffix}.xlsx")
        error = source.with_name(f"{source.stem}_{label}错误明细{suffix}.xlsx") if include_error else None
        if not result.exists() and (error is None or not error.exists()):
            return result, error
        index += 1


def choose_mixed_output_paths(
    source: Path,
    include_error: bool,
) -> tuple[Path, Path, Path | None]:
    """为混合模式的两个结果和错误明细分配相同的版本编号。"""
    index = 0
    while True:
        suffix = "" if index == 0 else f"({index})"
        shortlink = source.with_name(f"{source.stem}_生效短链导入结果{suffix}.xlsx")
        signature = source.with_name(f"{source.stem}_生效签名导入结果{suffix}.xlsx")
        error = source.with_name(f"{source.stem}_混合表错误明细{suffix}.xlsx") if include_error else None
        candidates = [shortlink, signature]
        if error is not None:
            candidates.append(error)
        if not any(path.exists() for path in candidates):
            return shortlink, signature, error
        index += 1


def choose_self_owned_output_paths(
    source: Path,
    include_error: bool,
) -> tuple[Path, Path | None]:
    index = 0
    while True:
        suffix = "" if index == 0 else f"({index})"
        shortlink = source.with_name(f"{source.stem}_生效短链导入结果{suffix}.xlsx")
        error = (
            source.with_name(f"{source.stem}_自家移动错误明细{suffix}.xlsx")
            if include_error
            else None
        )
        candidates = [shortlink]
        if error is not None:
            candidates.append(error)
        if not any(path.exists() for path in candidates):
            return shortlink, error
        index += 1


def choose_telecom_output_paths(
    batch_dir: Path,
    include_error: bool,
    include_failure_reason: bool,
) -> tuple[Path, Path, Path, Path, Path | None]:
    output_dir = batch_dir / TELECOM_OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    index = 0
    label = batch_dir.name or "自家电信批次"
    while True:
        suffix = "" if index == 0 else f"({index})"
        shortlink = output_dir / f"{label}_生效短链导入结果{suffix}.xlsx"
        signature = output_dir / f"{label}_生效签名导入结果{suffix}.xlsx"
        final_diversion = output_dir / f"最终引流报备表{suffix}.xlsx"
        failure_reason = output_dir / f"{label}_签名错误原因{suffix}.xlsx"
        error = (
            output_dir / f"{label}_自家电信错误明细{suffix}.xlsx"
            if include_error
            else None
        )
        candidates = [shortlink, signature, final_diversion]
        if include_failure_reason:
            candidates.append(failure_reason)
        if error is not None:
            candidates.append(error)
        if not any(path.exists() for path in candidates):
            return shortlink, signature, final_diversion, failure_reason, error
        index += 1


def telecom_input_fingerprint(paths: Sequence[Path]) -> dict[str, object]:
    entries: list[dict[str, object]] = []
    combined = hashlib.sha256()
    unique_paths = {path.resolve() for path in paths}
    for path in sorted(unique_paths, key=lambda item: str(item).lower()):
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        item_hash = digest.hexdigest()
        size = path.stat().st_size
        entries.append({"name": path.name, "size": size, "sha256": item_hash})
        combined.update(path.name.encode("utf-8"))
        combined.update(str(size).encode("ascii"))
        combined.update(item_hash.encode("ascii"))
    return {"batch_sha256": combined.hexdigest(), "files": entries}


def ensure_telecom_batch_not_reused(batch_dir: Path, fingerprint: dict[str, object]) -> None:
    manifest_paths = [
        batch_dir / TELECOM_OUTPUT_DIR / TELECOM_MANIFEST_FILE,
        batch_dir / TELECOM_MANIFEST_FILE,
    ]
    for manifest_path in manifest_paths:
        if not manifest_path.exists():
            continue
        try:
            previous = json.loads(manifest_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if previous.get("batch_sha256") == fingerprint.get("batch_sha256"):
            raise ProcessingError(
                f"该批次输入文件已经处理过。记录文件：{manifest_path}"
            )


def save_telecom_manifest(
    batch_dir: Path,
    fingerprint: dict[str, object],
    summary: ProcessSummary,
) -> None:
    content = {
        **fingerprint,
        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "summary": asdict(summary),
    }
    output_dir = batch_dir / TELECOM_OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / TELECOM_MANIFEST_FILE
    temporary_path = output_dir / f".{TELECOM_MANIFEST_FILE}.tmp"
    temporary_path.write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary_path, path)


def validate_parameters(
    source_path: Path,
    mode: str,
    channel: str,
    start_sub_port: int,
    operator: str,
    business_type: str,
    important_support: str,
    table_action: str = "",
    max_rows: int = 1000,
    signature_match_mode: str = SIGNATURE_MATCH_FULL_RULE,
) -> None:
    if mode not in {
        MODE_SHORTLINK,
        MODE_SIGNATURE,
        MODE_MIXED,
        MODE_SELF_MOBILE,
        MODE_SELF_TELECOM,
        MODE_TABLE_TOOLS,
        MODE_SIGNATURE_CONSISTENCY,
        MODE_PORT_ALLOCATION,
    }:
        raise ProcessingError("未知的导出类型。")
    if mode == MODE_TABLE_TOOLS:
        if table_action not in {TABLE_ACTION_SPLIT, TABLE_ACTION_MERGE}:
            raise ProcessingError("请选择“拆分总表”或“合并分表”。")
        if table_action == TABLE_ACTION_SPLIT:
            if not source_path.exists() or not source_path.is_file():
                raise ProcessingError("请选择存在的 XLSX/CSV 总表文件。")
            if source_path.suffix.lower() not in {".xlsx", ".csv"}:
                raise ProcessingError("拆分总表只支持 .xlsx 和 .csv 文件。")
            if max_rows < 2:
                raise ProcessingError("每个分表最大行数必须至少为 2（包含表头）。")
        elif not source_path.exists() or not source_path.is_dir():
            raise ProcessingError("请选择存在的分表文件夹。")
        return
    if mode == MODE_SELF_TELECOM:
        if not source_path.exists() or not source_path.is_dir():
            raise ProcessingError("请选择存在的自家电信批次文件夹。")
        if important_support not in {"是", "否"}:
            raise ProcessingError("请选择是否重保。")
        return
    if not source_path.exists() or not source_path.is_file():
        raise ProcessingError("请选择存在的客户源文件。")
    if mode == MODE_PORT_ALLOCATION:
        if source_path.suffix.lower() != ".xlsx":
            raise ProcessingError("子端口分配只支持 .xlsx 分配表。")
        return
    if mode == MODE_SIGNATURE_CONSISTENCY:
        if signature_match_mode not in {
            SIGNATURE_MATCH_FULL_RULE,
            SIGNATURE_MATCH_REGION_ONLY,
        }:
            raise ProcessingError("请选择“完整规则”或“只匹配地区”。")
        if source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            if source_path.suffix.lower() == ".xls":
                raise ProcessingError("不再支持旧版 .xls，请先在 Excel 中另存为 .xlsx 后再选择。")
            raise ProcessingError("签名一致性判定只支持 .xlsx 和 .xlsm 文件。")
        return
    if mode == MODE_SELF_MOBILE:
        if source_path.suffix.lower() != ".xlsx":
            if source_path.suffix.lower() == ".xls":
                raise ProcessingError("不再支持旧版 .xls，请先在 Excel 中另存为 .xlsx 后再选择。")
            raise ProcessingError("自家移动只支持 .xlsx 文件。")
        if important_support not in {"是", "否"}:
            raise ProcessingError("请选择是否重保。")
        return
    allowed = {".xlsx", ".xlsm", ".csv"}
    if mode == MODE_SIGNATURE:
        allowed.add(".txt")
    if source_path.suffix.lower() not in allowed:
        if source_path.suffix.lower() == ".xls":
            raise ProcessingError("不再支持旧版 .xls，请先在 Excel 中另存为 .xlsx 后再选择。")
        formats = "、".join(sorted(allowed))
        raise ProcessingError(f"当前模式只支持：{formats}")
    if not channel.strip():
        raise ProcessingError("请输入通道编号。")
    if start_sub_port < 0:
        raise ProcessingError("起始子端口必须是非负整数。")
    if operator not in {"移动", "联通", "电信"}:
        raise ProcessingError("请选择正确的运营商。")
    if business_type not in {"行业", "营销"}:
        raise ProcessingError("请选择正确的业务类型。")
    if important_support not in {"是", "否"}:
        raise ProcessingError("请选择是否重保。")


def process_file(
    source_path: Path,
    mode: str,
    channel: str,
    start_sub_port: int,
    operator: str,
    business_type: str,
    important_support: str,
    template_dir: Path | None = None,
    status_callback: Callable[[str], None] | None = None,
    table_action: str = "",
    max_rows: int = 1000,
    table_progress_callback: Callable[[str, int, int, str], None] | None = None,
    table_cancel_check: Callable[[], bool] | None = None,
    signature_match_mode: str = SIGNATURE_MATCH_FULL_RULE,
    master_path: Path | None = None,
    signature_total_dir: Path | None = None,
) -> ProcessSummary:
    """执行一次完整处理，是图形界面和命令行测试共同调用的入口。"""
    source_path = source_path.resolve()
    validate_parameters(
        source_path,
        mode,
        channel,
        start_sub_port,
        operator,
        business_type,
        important_support,
        table_action,
        max_rows,
        signature_match_mode,
    )
    if mode == MODE_PORT_ALLOCATION:
        if master_path is None:
            raise ProcessingError("请选择最新的历史总表。")
        if status_callback:
            status_callback("正在核对签名编号与历史子端口，并回填空缺……")
        try:
            result = allocate_ports(source_path, master_path)
        except PortAllocationError as exc:
            raise ProcessingError(str(exc)) from exc
        return ProcessSummary(
            output_path=result.allocation_path,
            secondary_output_path=result.master_path,
            error_path=None,
            generated_rows=result.allocated_rows,
            duplicate_rows=0,
            error_rows=0,
        )
    if mode == MODE_TABLE_TOOLS:
        try:
            if table_action == TABLE_ACTION_SPLIT:
                result = split_table(
                    source_path,
                    max_rows,
                    status_callback,
                    table_progress_callback,
                    table_cancel_check,
                )
            else:
                result = merge_tables(
                    source_path,
                    status_callback=status_callback,
                    progress_callback=table_progress_callback,
                    cancel_check=table_cancel_check,
                )
        except TableOperationError as exc:
            raise ProcessingError(str(exc)) from exc
        return ProcessSummary(
            output_path=result.output_path,
            error_path=None,
            generated_rows=result.generated_rows,
            duplicate_rows=0,
            error_rows=0,
            table_action=result.action,
            output_file_count=result.output_file_count,
            source_file_count=result.source_file_count,
            output_dir=result.output_dir,
            output_paths=result.output_paths,
        )
    if mode == MODE_SIGNATURE_CONSISTENCY:
        data_dir = (template_dir or application_dir()).resolve()
        region_library = data_dir / REGION_LIBRARY_FILE
        try:
            result = process_signature_consistency(
                source_path,
                region_library,
                status_callback=status_callback,
                match_mode=signature_match_mode,
            )
        except SignatureConsistencyError as exc:
            raise ProcessingError(str(exc)) from exc
        return ProcessSummary(
            output_path=result.output_path,
            error_path=None,
            generated_rows=result.source_rows,
            duplicate_rows=0,
            error_rows=0,
            skipped_rows=(
                result.first_layer_filtered_rows
                + result.second_layer_filtered_rows
                + result.enterprise_type_filtered_rows
            ),
            source_rows=result.source_rows,
            first_layer_rows=result.first_layer_rows,
            second_layer_rows=result.second_layer_rows,
            first_layer_filtered_rows=result.first_layer_filtered_rows,
            second_layer_filtered_rows=result.second_layer_filtered_rows,
            enterprise_type_filtered_rows=result.enterprise_type_filtered_rows,
            processed_sheets=result.processed_sheets,
        )
    if status_callback:
        status_callback("正在读取并整理客户数据……")

    if mode == MODE_SELF_MOBILE:
        data_dir = (template_dir or application_dir()).resolve()
        shortlink_template = data_dir / SHORTLINK_TEMPLATE
        if not shortlink_template.exists():
            raise ProcessingError(
                f"内置模板缺失：{SHORTLINK_TEMPLATE}。请重新安装程序。"
            )
        (
            shortlink_records,
            errors,
            duplicate_count,
            skipped_count,
            dirty_count,
        ) = read_self_owned_records(
            source_path,
            "移动",
            important_support,
            data_dir,
            signature_total_dir,
        )
        shortlink_output, error_path = choose_self_owned_output_paths(
            source_path,
            bool(errors),
        )
        if error_path is not None:
            if status_callback:
                status_callback("正在生成自家移动错误明细……")
            save_self_owned_error_report(error_path, errors)
        if not shortlink_records:
            detail = f"错误明细已保存到：{error_path}" if error_path else ""
            dirty_detail = (
                f"已忽略 {dirty_count:,} 条不属于本公司主码号的数据。" if dirty_count else ""
            )
            raise ProcessingError(f"没有读取到可导出的已生效引流数据。{dirty_detail}{detail}")

        if status_callback:
            status_callback(f"正在生成自家移动引流，共 {len(shortlink_records):,} 条……")
        write_self_owned_output(
            shortlink_template,
            shortlink_output,
            MODE_SHORTLINK,
            shortlink_records,
        )
        generated_shortlink_path = str(shortlink_output)
        return ProcessSummary(
            output_path=generated_shortlink_path,
            error_path=str(error_path) if error_path else None,
            generated_rows=len(shortlink_records),
            duplicate_rows=duplicate_count,
            error_rows=len(errors),
            shortlink_rows=len(shortlink_records),
            signature_rows=0,
            shortlink_output_path=generated_shortlink_path,
            skipped_rows=skipped_count,
            dirty_rows=dirty_count,
        )

    if mode == MODE_SELF_TELECOM:
        data_dir = (template_dir or application_dir()).resolve()
        shortlink_template = data_dir / SHORTLINK_TEMPLATE
        signature_template = data_dir / SIGNATURE_TEMPLATE
        final_diversion_template = data_dir / FINAL_DIVERSION_TEMPLATE
        missing_templates = [
            path.name
            for path in (shortlink_template, signature_template, final_diversion_template)
            if not path.exists()
        ]
        if missing_templates:
            raise ProcessingError("未找到模板：" + "、".join(missing_templates))
        (
            shortlink_records,
            signature_records,
            final_diversion_records,
            errors,
            failure_reason_records,
            duplicate_count,
            failed_count,
            batch_count,
            unrelated_failed_count,
            dirty_count,
            input_paths,
        ) = read_self_telecom_records(source_path, important_support, data_dir)
        fingerprint = telecom_input_fingerprint(
            [
                *input_paths,
                shortlink_template,
                signature_template,
                final_diversion_template,
            ]
        )
        ensure_telecom_batch_not_reused(source_path, fingerprint)
        (
            shortlink_output,
            signature_output,
            final_diversion_output,
            failure_reason_output,
            error_path,
        ) = choose_telecom_output_paths(
            source_path,
            bool(errors),
            bool(failure_reason_records),
        )
        generated_failure_reason_path: str | None = None
        if failure_reason_records:
            if status_callback:
                status_callback(
                    f"正在生成签名错误原因，共 {len(failure_reason_records):,} 条……"
                )
            save_telecom_failure_reason_report(
                failure_reason_output,
                failure_reason_records,
            )
            generated_failure_reason_path = str(failure_reason_output)
        if error_path is not None:
            if status_callback:
                status_callback("正在生成自家电信错误明细……")
            save_telecom_error_report(error_path, errors)
        if not shortlink_records and not signature_records:
            details: list[str] = []
            if generated_failure_reason_path:
                details.append(f"签名错误原因已保存到：{generated_failure_reason_path}")
            if error_path:
                details.append(f"错误明细已保存到：{error_path}")
            detail = "；".join(details)
            raise ProcessingError(f"没有可导出的电信成功数据。{detail}")

        generated_shortlink_path: str | None = None
        generated_signature_path: str | None = None
        generated_final_diversion_path: str | None = None
        if shortlink_records:
            if status_callback:
                status_callback(
                    f"正在生成最终引流报备表，共 {len(final_diversion_records):,} 条……"
                )
            write_final_diversion_output(
                final_diversion_template,
                final_diversion_output,
                final_diversion_records,
            )
            generated_final_diversion_path = str(final_diversion_output)
            if status_callback:
                status_callback(f"正在生成自家电信短链，共 {len(shortlink_records):,} 条……")
            write_self_owned_output(
                shortlink_template, shortlink_output, MODE_SHORTLINK, shortlink_records
            )
            generated_shortlink_path = str(shortlink_output)
        if signature_records:
            if status_callback:
                status_callback(f"正在生成自家电信签名，共 {len(signature_records):,} 条……")
            write_self_owned_output(
                signature_template, signature_output, MODE_SIGNATURE, signature_records
            )
            generated_signature_path = str(signature_output)
        summary = ProcessSummary(
            output_path=generated_shortlink_path or generated_signature_path,
            secondary_output_path=(
                generated_signature_path
                if generated_shortlink_path and generated_signature_path
                else None
            ),
            error_path=str(error_path) if error_path else None,
            generated_rows=len(shortlink_records) + len(signature_records),
            duplicate_rows=duplicate_count,
            error_rows=len(errors),
            shortlink_rows=len(shortlink_records),
            signature_rows=len(signature_records),
            shortlink_output_path=generated_shortlink_path,
            signature_output_path=generated_signature_path,
            failed_rows=failed_count,
            batch_count=batch_count,
            dirty_rows=dirty_count,
            unrelated_failed_rows=unrelated_failed_count,
            final_diversion_rows=len(final_diversion_records),
            final_diversion_output_path=generated_final_diversion_path,
            failure_reason_rows=len(failure_reason_records),
            failure_reason_output_path=generated_failure_reason_path,
        )
        save_telecom_manifest(source_path, fingerprint, summary)
        return summary

    if mode == MODE_MIXED:
        (
            shortlink_records,
            signature_records,
            errors,
            shortlink_duplicate_count,
            signature_duplicate_count,
        ) = read_mixed_records(source_path)
        templates = (template_dir or application_dir()).resolve()
        shortlink_template = templates / SHORTLINK_TEMPLATE
        signature_template = templates / SIGNATURE_TEMPLATE
        missing_templates = [
            path.name for path in (shortlink_template, signature_template) if not path.exists()
        ]
        if missing_templates:
            raise ProcessingError(
                "未找到模板："
                + "、".join(missing_templates)
                + "。安装版请重新安装程序；源码运行请指定模板目录。"
            )

        shortlink_output, signature_output, error_path = choose_mixed_output_paths(
            source_path,
            bool(errors),
        )
        if error_path is not None:
            if status_callback:
                status_callback("正在生成混合表错误明细……")
            save_error_report(error_path, errors, include_category=True)

        if not shortlink_records and not signature_records:
            detail = f"错误明细已保存到：{error_path}" if error_path else ""
            raise ProcessingError(f"没有读取到可导出的有效数据。{detail}")

        generated_shortlink_path: str | None = None
        generated_signature_path: str | None = None
        # 先生成短链并占用起始子端口；签名从短链末尾继续编号。
        if shortlink_records:
            if status_callback:
                status_callback(f"正在生成生效短链，共 {len(shortlink_records):,} 条……")
            write_template_output(
                shortlink_template,
                shortlink_output,
                MODE_SHORTLINK,
                shortlink_records,
                channel.strip(),
                start_sub_port,
                operator,
                business_type,
                important_support,
            )
            generated_shortlink_path = str(shortlink_output)
        if signature_records:
            if status_callback:
                status_callback(f"正在生成生效签名，共 {len(signature_records):,} 条……")
            write_template_output(
                signature_template,
                signature_output,
                MODE_SIGNATURE,
                signature_records,
                channel.strip(),
                start_sub_port + len(shortlink_records),
                operator,
                business_type,
                important_support,
            )
            generated_signature_path = str(signature_output)

        primary_output = generated_shortlink_path or generated_signature_path
        secondary_output = (
            generated_signature_path
            if generated_shortlink_path and generated_signature_path
            else None
        )
        return ProcessSummary(
            output_path=primary_output,
            error_path=str(error_path) if error_path else None,
            generated_rows=len(shortlink_records) + len(signature_records),
            duplicate_rows=shortlink_duplicate_count + signature_duplicate_count,
            error_rows=len(errors),
            secondary_output_path=secondary_output,
            shortlink_rows=len(shortlink_records),
            signature_rows=len(signature_records),
            shortlink_output_path=generated_shortlink_path,
            signature_output_path=generated_signature_path,
        )

    if mode == MODE_SHORTLINK:
        records, errors, duplicate_count = read_shortlink_records(source_path)
        template_name = SHORTLINK_TEMPLATE
    elif mode == MODE_SIGNATURE:
        records, errors, duplicate_count = read_signature_records(source_path)
        template_name = SIGNATURE_TEMPLATE
    else:
        raise ProcessingError("未知的导出类型。")

    output_path, error_path = choose_output_paths(source_path, mode, bool(errors))
    if error_path is not None:
        if status_callback:
            status_callback("正在生成错误明细……")
        save_error_report(error_path, errors)

    if not records:
        detail = f"错误明细已保存到：{error_path}" if error_path else ""
        raise ProcessingError(f"没有读取到可导出的有效数据。{detail}")

    templates = (template_dir or application_dir()).resolve()
    template_path = templates / template_name
    if not template_path.exists():
        raise ProcessingError(
            f"未找到模板“{template_name}”。安装版请重新安装程序；源码运行请指定模板目录。"
        )

    if status_callback:
        status_callback(f"正在写入模板，共 {len(records):,} 条……")
    write_template_output(
        template_path,
        output_path,
        mode,
        records,
        channel.strip(),
        start_sub_port,
        operator,
        business_type,
        important_support,
    )
    return ProcessSummary(
        output_path=str(output_path),
        error_path=str(error_path) if error_path else None,
        generated_rows=len(records),
        duplicate_rows=duplicate_count,
        error_rows=len(errors),
        shortlink_rows=len(records) if mode == MODE_SHORTLINK else 0,
        signature_rows=len(records) if mode == MODE_SIGNATURE else 0,
        shortlink_output_path=str(output_path) if mode == MODE_SHORTLINK else None,
        signature_output_path=str(output_path) if mode == MODE_SIGNATURE else None,
    )


class ImporterApp(tk.Tk):
    """面向不使用 Python 的同事提供的 Windows 图形界面。"""
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_NAME)
        self.geometry("1040x720")
        self.minsize(980, 670)
        self.configure(bg="#F3F6FA")
        self._worker: threading.Thread | None = None

        self.mode = tk.StringVar(value=MODE_SHORTLINK)
        self.source_path = tk.StringVar()
        self.master_path = tk.StringVar()
        self.channel = tk.StringVar()
        self.start_sub_port = tk.StringVar(value="1000")
        self.operator = tk.StringVar(value="移动")
        self.business_type = tk.StringVar(value="行业")
        self.important_support = tk.StringVar(value="否")
        self.table_action = tk.StringVar(value="拆分总表")
        self.signature_match_mode = tk.StringVar(value="完整规则")
        self.max_rows = tk.StringVar(value="1000")
        self.status = tk.StringVar(value="请选择客户源文件并填写参数")

        self._configure_styles()
        self._build_ui()
        self.mode.trace_add("write", self._mode_changed)
        self.table_action.trace_add("write", self._table_action_changed)
        self.signature_match_mode.trace_add("write", self._signature_match_mode_changed)
        self._update_parameter_states()
        self._refresh_template_status()

    def _configure_styles(self) -> None:
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("TFrame", background="#F3F6FA")
        style.configure("Header.TFrame", background="#132B46")
        style.configure("Card.TFrame", background="#FFFFFF")
        style.configure("TLabel", background="#F3F6FA", font=("Microsoft YaHei UI", 10))
        style.configure("Card.TLabel", background="#FFFFFF", font=("Microsoft YaHei UI", 10))
        style.configure(
            "Title.TLabel",
            background="#132B46",
            foreground="#FFFFFF",
            font=("Microsoft YaHei UI", 20, "bold"),
        )
        style.configure(
            "Subtitle.TLabel",
            background="#132B46",
            foreground="#B4CAE0",
            font=("Microsoft YaHei UI", 9),
        )
        style.configure(
            "Mode.TRadiobutton",
            background="#FFFFFF",
            foreground="#26405B",
            font=("Microsoft YaHei UI", 10),
        )
        style.configure(
            "Primary.TButton",
            background="#0F739A",
            foreground="#FFFFFF",
            font=("Microsoft YaHei UI", 11, "bold"),
            padding=(24, 12),
        )
        style.map("Primary.TButton", background=[("active", "#0C5F80"), ("disabled", "#9CB6CD")])
        style.configure("TButton", font=("Microsoft YaHei UI", 9), padding=(10, 7))
        style.configure("TEntry", padding=6)
        style.configure("TCombobox", padding=5)

    def _build_ui(self) -> None:
        outer = ttk.Frame(self, padding=(28, 20))
        outer.pack(fill="both", expand=True)
        header = ttk.Frame(outer, padding=(25, 18), style="Header.TFrame")
        header.pack(fill="x", pady=(0, 16))
        ttk.Label(header, text=APP_NAME, style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            header,
            text="通道处理  ·  表格整理  ·  子端口分配",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))

        card = ttk.Frame(outer, style="Card.TFrame", padding=(26, 20))
        card.pack(fill="both", expand=True)
        card.columnconfigure(1, weight=1)

        ttk.Label(card, text="导出类型", style="Card.TLabel").grid(row=0, column=0, sticky="w", pady=7)
        mode_frame = ttk.Frame(card, style="Card.TFrame")
        mode_frame.grid(row=0, column=1, sticky="w", pady=7)
        modes = (
            ("生效短链", MODE_SHORTLINK),
            ("生效签名", MODE_SIGNATURE),
            ("混合表自动分流", MODE_MIXED),
            ("自家移动", MODE_SELF_MOBILE),
            ("自家电信", MODE_SELF_TELECOM),
            ("拆分/合并表", MODE_TABLE_TOOLS),
            ("签名一致性判定", MODE_SIGNATURE_CONSISTENCY),
            ("子端口分配", MODE_PORT_ALLOCATION),
        )
        for index, (label, value) in enumerate(modes):
            ttk.Radiobutton(
                mode_frame,
                text=label,
                value=value,
                variable=self.mode,
                style="Mode.TRadiobutton",
            ).grid(row=index // 3, column=index % 3, sticky="w", padx=(0, 28), pady=3)

        self.source_label = ttk.Label(card, text="客户源文件", style="Card.TLabel")
        self.source_label.grid(row=1, column=0, sticky="w", pady=7)
        file_frame = ttk.Frame(card, style="Card.TFrame")
        file_frame.grid(row=1, column=1, sticky="ew", pady=7)
        file_frame.columnconfigure(0, weight=1)
        self.file_entry = ttk.Entry(file_frame, textvariable=self.source_path, state="readonly")
        self.file_entry.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.browse_button = ttk.Button(file_frame, text="选择文件", command=self._browse)
        self.browse_button.grid(row=0, column=1)

        self.master_label = ttk.Label(card, text="历史总表", style="Card.TLabel")
        self.master_label.grid(row=2, column=0, sticky="w", pady=7)
        master_frame = ttk.Frame(card, style="Card.TFrame")
        master_frame.grid(row=2, column=1, sticky="ew", pady=7)
        master_frame.columnconfigure(0, weight=1)
        self.master_entry = ttk.Entry(master_frame, textvariable=self.master_path, state="readonly")
        self.master_entry.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.master_browse_button = ttk.Button(master_frame, text="选择总表", command=self._browse_master)
        self.master_browse_button.grid(row=0, column=1)
        self.master_frame = master_frame

        self.channel_label = ttk.Label(card, text="通道编号", style="Card.TLabel")
        self.channel_label.grid(row=2, column=0, sticky="w", pady=7)
        self.channel_entry = ttk.Entry(card, textvariable=self.channel)
        self.channel_entry.grid(row=2, column=1, sticky="ew", pady=7)

        self.port_label = ttk.Label(card, text="起始子端口", style="Card.TLabel")
        self.port_label.grid(row=3, column=0, sticky="w", pady=7)
        self.port_entry = ttk.Entry(card, textvariable=self.start_sub_port)
        self.port_entry.grid(row=3, column=1, sticky="ew", pady=7)

        self.operator_label = ttk.Label(card, text="运营商", style="Card.TLabel")
        self.operator_label.grid(row=4, column=0, sticky="w", pady=7)
        self.operator_combo = ttk.Combobox(
            card, textvariable=self.operator, values=("移动", "联通", "电信"), state="readonly"
        )
        self.operator_combo.grid(row=4, column=1, sticky="ew", pady=7)

        self.business_label = ttk.Label(card, text="业务类型", style="Card.TLabel")
        self.business_label.grid(row=5, column=0, sticky="w", pady=7)
        self.business_combo = ttk.Combobox(
            card, textvariable=self.business_type, values=("行业", "营销"), state="readonly"
        )
        self.business_combo.grid(row=5, column=1, sticky="ew", pady=7)

        self.support_label = ttk.Label(card, text="是否重保", style="Card.TLabel")
        self.support_label.grid(row=6, column=0, sticky="w", pady=7)
        self.support_combo = ttk.Combobox(
            card, textvariable=self.important_support, values=("否", "是"), state="readonly"
        )
        self.support_combo.grid(row=6, column=1, sticky="ew", pady=7)

        self.table_action_label = ttk.Label(card, text="操作方式", style="Card.TLabel")
        self.table_action_label.grid(row=2, column=0, sticky="w", pady=7)
        self.table_action_combo = ttk.Combobox(
            card,
            textvariable=self.table_action,
            values=("拆分总表", "合并分表"),
            state="readonly",
        )
        self.table_action_combo.grid(row=2, column=1, sticky="ew", pady=7)
        self.signature_match_mode_label = ttk.Label(card, text="判定规则", style="Card.TLabel")
        self.signature_match_mode_label.grid(row=2, column=0, sticky="w", pady=7)
        self.signature_match_mode_combo = ttk.Combobox(
            card,
            textvariable=self.signature_match_mode,
            values=("完整规则", "只匹配地区"),
            state="readonly",
        )
        self.signature_match_mode_combo.grid(row=2, column=1, sticky="ew", pady=7)
        self.max_rows_label = ttk.Label(card, text="每个分表最大行数", style="Card.TLabel")
        self.max_rows_label.grid(row=3, column=0, sticky="w", pady=7)
        self.max_rows_entry = ttk.Entry(card, textvariable=self.max_rows)
        self.max_rows_entry.grid(row=3, column=1, sticky="ew", pady=7)

        self.template_status = tk.Label(
            card,
            text="",
            bg="#FFFFFF",
            fg="#4B647A",
            font=("Microsoft YaHei UI", 9),
            anchor="w",
        )
        self.template_status.grid(row=7, column=0, columnspan=2, sticky="ew", pady=(12, 4))

        self.progress = ttk.Progressbar(card, mode="indeterminate")
        self.progress.grid(row=8, column=0, columnspan=2, sticky="ew", pady=(8, 6))

        self.status_label = tk.Label(
            card,
            textvariable=self.status,
            bg="#FFFFFF",
            fg="#62748A",
            font=("Microsoft YaHei UI", 9),
            anchor="w",
        )
        self.status_label.grid(row=9, column=0, columnspan=2, sticky="ew", pady=(0, 10))

        self.run_button = ttk.Button(
            card, text="开始生成", style="Primary.TButton", command=self._start
        )
        self.run_button.grid(row=10, column=0, columnspan=2, pady=(5, 0))

        ttk.Label(
            outer,
            text="结果默认保存在客户源文件所在文件夹；只有存在异常数据时才生成错误明细。",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(12, 0))

    def _mode_changed(self, *_args: object) -> None:
        selected = self.source_path.get()
        if self.mode.get() != MODE_SIGNATURE and selected.lower().endswith(".txt"):
            self.source_path.set("")
        if self.mode.get() == MODE_SELF_MOBILE and selected and not selected.lower().endswith(".xlsx"):
            self.source_path.set("")
        if self.mode.get() == MODE_SELF_TELECOM and selected and not Path(selected).is_dir():
            self.source_path.set("")
        if self.mode.get() == MODE_TABLE_TOOLS and selected:
            expected_directory = self.table_action.get() == "合并分表"
            if Path(selected).is_dir() != expected_directory:
                self.source_path.set("")
        self._update_parameter_states()
        self._refresh_template_status()

    def _table_action_changed(self, *_args: object) -> None:
        if self.mode.get() != MODE_TABLE_TOOLS:
            return
        self.source_path.set("")
        self._update_parameter_states()

    def _table_action_value(self) -> str:
        return TABLE_ACTION_MERGE if self.table_action.get() == "合并分表" else TABLE_ACTION_SPLIT

    def _signature_match_mode_changed(self, *_args: object) -> None:
        if self.mode.get() == MODE_SIGNATURE_CONSISTENCY:
            self._update_parameter_states()

    def _signature_match_mode_value(self) -> str:
        return (
            SIGNATURE_MATCH_REGION_ONLY
            if self.signature_match_mode.get() == "只匹配地区"
            else SIGNATURE_MATCH_FULL_RULE
        )

    def _update_parameter_states(self) -> None:
        self_owned = self.mode.get() in {MODE_SELF_MOBILE, MODE_SELF_TELECOM}
        table_tools = self.mode.get() == MODE_TABLE_TOOLS
        consistency = self.mode.get() == MODE_SIGNATURE_CONSISTENCY
        port_allocation = self.mode.get() == MODE_PORT_ALLOCATION
        merging = table_tools and self._table_action_value() == TABLE_ACTION_MERGE
        if table_tools:
            self.source_label.configure(text="分表文件夹" if merging else "总表文件")
            self.browse_button.configure(text="选择文件夹" if merging else "选择文件")
        elif port_allocation:
            self.source_label.configure(text="待分配表")
            self.browse_button.configure(text="选择分配表")
        elif consistency:
            self.source_label.configure(text="待判定文件")
            self.browse_button.configure(text="选择文件")
        else:
            self.source_label.configure(
                text="批次文件夹" if self.mode.get() == MODE_SELF_TELECOM else "客户源文件"
            )
            self.browse_button.configure(
                text="选择文件夹" if self.mode.get() == MODE_SELF_TELECOM else "选择文件"
            )
        parameter_widgets = (
            self.channel_label, self.channel_entry,
            self.port_label, self.port_entry,
            self.operator_label, self.operator_combo,
            self.business_label, self.business_combo,
        )
        for widget in parameter_widgets:
            if self_owned or table_tools or consistency or port_allocation:
                widget.grid_remove()
            else:
                widget.grid()
        for widget in (self.support_label, self.support_combo):
            if table_tools or consistency or port_allocation:
                widget.grid_remove()
            else:
                widget.grid()
        for widget in (self.table_action_label, self.table_action_combo):
            if table_tools:
                widget.grid()
            else:
                widget.grid_remove()
        for widget in (self.signature_match_mode_label, self.signature_match_mode_combo):
            if consistency:
                widget.grid()
            else:
                widget.grid_remove()
        for widget in (self.max_rows_label, self.max_rows_entry):
            if table_tools and not merging:
                widget.grid()
            else:
                widget.grid_remove()
        if port_allocation:
            self.master_label.grid()
            self.master_frame.grid()
        else:
            self.master_label.grid_remove()
            self.master_frame.grid_remove()
        if table_tools:
            self.run_button.configure(text="开始合并" if merging else "开始拆分")
            self.status.set(
                "请选择包含分表的文件夹" if merging else "请选择总表并设置每个分表最大行数"
            )
        elif consistency:
            self.run_button.configure(text="开始判定")
            self.status.set(
                "只校验签名与企业名称是否含同一地区"
                if self._signature_match_mode_value() == SIGNATURE_MATCH_REGION_ONLY
                else "按顺序跳字并校验地区，保留全部原数据"
            )
        elif port_allocation:
            self.run_button.configure(text="开始分配")
            self.status.set("选择待分配表和最新历史总表；结果保存为两份新文件")
        else:
            self.run_button.configure(text="开始生成")

    def _refresh_template_status(self) -> None:
        if self.mode.get() == MODE_PORT_ALLOCATION:
            self.template_status.configure(
                text="行业 701–999 · 营销 001–699 · 按主码号和签名编号查找空号",
                fg="#14804A",
            )
            return
        if self.mode.get() == MODE_TABLE_TOOLS:
            action = "合并文件夹内表头完全一致的分表" if self._table_action_value() == TABLE_ACTION_MERGE else "按包含表头的最大行数拆分总表"
            self.template_status.configure(
                text=f"表格工具：{action}；结果保存到“生成结果”文件夹",
                fg="#14804A",
            )
            return
        if self.mode.get() == MODE_SIGNATURE_CONSISTENCY:
            region_path = application_dir() / REGION_LIBRARY_FILE
            self.template_status.configure(
                text=(
                    f"地区库：{REGION_LIBRARY_FILE}（"
                    f"{'已找到' if region_path.exists() else '未找到，请与程序放在同一文件夹'}）"
                ),
                fg="#14804A" if region_path.exists() else "#B42318",
            )
            return
        if self.mode.get() == MODE_SELF_MOBILE:
            names = (SHORTLINK_TEMPLATE, CHANNEL_MAPPING_FILE)
        elif self.mode.get() == MODE_SELF_TELECOM:
            names = (
                SHORTLINK_TEMPLATE,
                SIGNATURE_TEMPLATE,
                FINAL_DIVERSION_TEMPLATE,
                CHANNEL_MAPPING_FILE,
            )
        elif self.mode.get() == MODE_MIXED:
            names = (SHORTLINK_TEMPLATE, SIGNATURE_TEMPLATE)
        else:
            names = (
                SHORTLINK_TEMPLATE if self.mode.get() == MODE_SHORTLINK else SIGNATURE_TEMPLATE,
            )
        missing = [name for name in names if not (application_dir() / name).exists()]
        found = not missing
        name_text = "、".join(names)
        label = (
            "所需文件"
            if self.mode.get() in {MODE_SELF_MOBILE, MODE_SELF_TELECOM}
            else "模板"
        )
        self.template_status.configure(
            text=f"{label}：{name_text}（{'已找到' if found else '未找到，请与程序放在同一文件夹'}）",
            fg="#14804A" if found else "#B42318",
        )

    def _browse(self) -> None:
        if self.mode.get() == MODE_PORT_ALLOCATION:
            selected = filedialog.askopenfilename(
                title="选择待分配表", filetypes=[("Excel 文件", "*.xlsx")]
            )
            if selected:
                self.source_path.set(selected)
                self.status.set("待分配表已选择；请继续选择最新历史总表")
            return
        if self.mode.get() == MODE_TABLE_TOOLS:
            if self._table_action_value() == TABLE_ACTION_MERGE:
                selected = filedialog.askdirectory(title="选择包含全部待合并分表的文件夹")
                if selected:
                    self.source_path.set(selected)
                    self.status.set("分表文件夹已选择，可以开始合并")
                return
            selected = filedialog.askopenfilename(
                title="选择需要拆分的总表",
                filetypes=[
                    ("支持的文件", "*.xlsx *.csv"),
                    ("Excel 文件", "*.xlsx"),
                    ("CSV 文件", "*.csv"),
                ],
            )
            if selected:
                self.source_path.set(selected)
                self.status.set("总表已选择，可以开始拆分")
            return
        if self.mode.get() == MODE_SELF_TELECOM:
            selected = filedialog.askdirectory(title="选择自家电信批次文件夹")
            if selected:
                self.source_path.set(selected)
                self.status.set("批次文件夹已选择，可以开始生成")
            return
        if self.mode.get() == MODE_SELF_MOBILE:
            filetypes = [
                ("运营商报备表", "*.xlsx"),
                ("Excel 文件", "*.xlsx"),
                ("所有文件", "*.*"),
            ]
        elif self.mode.get() == MODE_SIGNATURE:
            filetypes = [
                ("支持的文件", "*.xlsx *.xlsm *.csv *.txt"),
                ("Excel 文件", "*.xlsx *.xlsm"),
                ("CSV 文件", "*.csv"),
                ("文本文件", "*.txt"),
                ("所有文件", "*.*"),
            ]
        elif self.mode.get() == MODE_SIGNATURE_CONSISTENCY:
            filetypes = [
                ("Excel 文件", "*.xlsx *.xlsm"),
                ("XLSX 文件", "*.xlsx"),
                ("所有文件", "*.*"),
            ]
        else:
            filetypes = [
                ("支持的文件", "*.xlsx *.xlsm *.csv"),
                ("Excel 文件", "*.xlsx *.xlsm"),
                ("CSV 文件", "*.csv"),
                ("所有文件", "*.*"),
            ]
        selected = filedialog.askopenfilename(title="选择客户源文件", filetypes=filetypes)
        if selected:
            self.source_path.set(selected)
            self.status.set("文件已选择，可以开始生成")

    def _browse_master(self) -> None:
        selected = filedialog.askopenfilename(
            title="选择最新历史总表", filetypes=[("Excel 文件", "*.xlsx")]
        )
        if selected:
            self.master_path.set(selected)
            self.status.set("两个文件已选择，可以开始分配")

    def _set_busy(self, busy: bool) -> None:
        state = "disabled" if busy else "normal"
        self.run_button.configure(state=state)
        self.browse_button.configure(state=state)
        self.master_browse_button.configure(state=state)
        self.channel_entry.configure(state=state)
        self.port_entry.configure(state=state)
        combo_state = "disabled" if busy else "readonly"
        self.operator_combo.configure(state=combo_state)
        self.business_combo.configure(state=combo_state)
        self.support_combo.configure(state=combo_state)
        self.table_action_combo.configure(state=combo_state)
        self.signature_match_mode_combo.configure(state=combo_state)
        self.max_rows_entry.configure(state=state)
        if busy:
            self.progress.start(12)
        else:
            self.progress.stop()
            self._update_parameter_states()

    def _start(self) -> None:
        table_tools = self.mode.get() == MODE_TABLE_TOOLS
        table_action = self._table_action_value() if table_tools else ""
        max_rows = 1000
        if table_tools:
            start_port = 0
            if table_action == TABLE_ACTION_SPLIT:
                try:
                    max_rows = int(self.max_rows.get().strip())
                except ValueError:
                    messagebox.showerror(APP_NAME, "每个分表最大行数必须是整数。")
                    return
                if max_rows < 2:
                    messagebox.showerror(APP_NAME, "每个分表最大行数必须至少为 2（包含表头）。")
                    return
        elif self.mode.get() in {
            MODE_SELF_MOBILE,
            MODE_SELF_TELECOM,
            MODE_SIGNATURE_CONSISTENCY,
            MODE_PORT_ALLOCATION,
        }:
            start_port = 0
        else:
            try:
                start_port = int(self.start_sub_port.get().strip())
            except ValueError:
                messagebox.showerror(APP_NAME, "起始子端口必须是整数。")
                return
        if not self.source_path.get().strip():
            messagebox.showerror(
                APP_NAME,
                "请先选择分表文件夹。"
                if table_action == TABLE_ACTION_MERGE
                else "请先选择总表文件。"
                if table_tools
                else "请先选择客户源文件。",
            )
            return
        if self.mode.get() == MODE_PORT_ALLOCATION and not self.master_path.get().strip():
            messagebox.showerror(APP_NAME, "请先选择最新的历史总表。")
            return

        self._set_busy(True)
        self.status.set("正在开始处理……")
        arguments = {
            "source_path": Path(self.source_path.get()),
            "mode": self.mode.get(),
            "channel": self.channel.get(),
            "start_sub_port": start_port,
            "operator": (
                "移动" if self.mode.get() == MODE_SELF_MOBILE
                else "电信" if self.mode.get() == MODE_SELF_TELECOM
                else self.operator.get()
            ),
            "business_type": self.business_type.get(),
            "important_support": self.important_support.get(),
            "status_callback": self._thread_status,
            "table_action": table_action,
            "max_rows": max_rows,
            "signature_match_mode": self._signature_match_mode_value(),
            "master_path": Path(self.master_path.get()) if self.mode.get() == MODE_PORT_ALLOCATION else None,
        }
        self._worker = threading.Thread(target=self._run_worker, kwargs=arguments, daemon=True)
        self._worker.start()

    def _thread_status(self, text: str) -> None:
        self.after(0, self.status.set, text)

    def _run_worker(self, **kwargs: object) -> None:
        try:
            summary = process_file(**kwargs)
        except Exception as exc:
            self.after(0, self._finish_error, exc)
        else:
            self.after(0, self._finish_success, summary)

    def _finish_error(self, exc: Exception) -> None:
        self._set_busy(False)
        self.status.set("处理失败，请根据提示检查")
        message = str(exc) if isinstance(exc, ProcessingError) else f"发生未预期错误：{exc}"
        messagebox.showerror(APP_NAME, message)

    def _finish_success(self, summary: ProcessSummary) -> None:
        self._set_busy(False)
        self.status.set(f"处理完成：已生成 {summary.generated_rows:,} 条")
        if self.mode.get() == MODE_TABLE_TOOLS:
            if summary.table_action == TABLE_ACTION_SPLIT:
                lines = [
                    "总表拆分完成。",
                    f"数据行数：{summary.generated_rows:,} 条",
                    f"生成分表：{summary.output_file_count:,} 个",
                    "",
                    f"结果文件夹：\n{summary.output_dir}",
                ]
            else:
                lines = [
                    "分表合并完成。",
                    f"读取分表：{summary.source_file_count:,} 个",
                    f"合并数据：{summary.generated_rows:,} 条",
                    "",
                    f"合并结果：\n{summary.output_path}",
                ]
        elif self.mode.get() == MODE_PORT_ALLOCATION:
            lines = [
                "子端口分配完成。",
                f"已分配：{summary.generated_rows:,} 行（每行各有行业、营销一个号码）",
                "",
                f"分配结果：\n{summary.output_path}",
                f"更新总表：\n{summary.secondary_output_path}",
                "下次分配请使用这份更新总表。",
            ]
        elif self.mode.get() == MODE_SIGNATURE_CONSISTENCY:
            lines = [
                "签名一致性判定完成。",
                f"读取数据：{summary.source_rows:,} 条",
                f"第一层通过：{summary.first_layer_rows:,} 条",
                f"第一层过滤：{summary.first_layer_filtered_rows:,} 条",
                f"第二层通过：{summary.second_layer_rows:,} 条",
                f"第二层过滤：{summary.second_layer_filtered_rows:,} 条",
                f"企业类型判定失败：{summary.enterprise_type_filtered_rows:,} 条",
                f"处理工作表：{summary.processed_sheets:,} 个",
                "",
                f"结果文件：\n{summary.output_path}",
            ]
        elif self.mode.get() == MODE_SELF_MOBILE:
            lines = [
                "自家移动导入完成。",
                f"生效引流：{summary.shortlink_rows:,} 条",
                f"完全重复并去重：{summary.duplicate_rows:,} 条",
                f"异常并跳过：{summary.error_rows:,} 条",
                f"非已生效并跳过：{summary.skipped_rows:,} 条",
                f"非本公司主码号并忽略：{summary.dirty_rows:,} 条",
                "",
                f"生效引流结果：\n{summary.shortlink_output_path}",
            ]
        elif self.mode.get() == MODE_SELF_TELECOM:
            lines = [
                "自家电信批次导入完成。",
                f"已配对提交批次：{summary.batch_count:,} 组",
                f"失败回执剔除：{summary.failed_rows:,} 条",
                f"签名错误原因：{summary.failure_reason_rows:,} 条",
                f"生效短链：{summary.shortlink_rows:,} 条",
                f"生效签名：{summary.signature_rows:,} 条",
                f"最终引流报备：{summary.final_diversion_rows:,} 条",
                f"完全重复并去重：{summary.duplicate_rows:,} 条",
                f"异常并跳过：{summary.error_rows:,} 条",
            ]
            if summary.dirty_rows:
                lines.append(f"非本公司电信主码号并忽略：{summary.dirty_rows:,} 条")
            if summary.unrelated_failed_rows:
                lines.append(
                    f"回执中非本批次子端口并忽略：{summary.unrelated_failed_rows:,} 条"
                )
            if summary.shortlink_output_path:
                lines.extend(["", f"生效短链结果：\n{summary.shortlink_output_path}"])
            if summary.signature_output_path:
                lines.extend(["", f"生效签名结果：\n{summary.signature_output_path}"])
            if summary.final_diversion_output_path:
                lines.extend(["", f"最终引流报备表：\n{summary.final_diversion_output_path}"])
            if summary.failure_reason_output_path:
                lines.extend(["", f"签名错误原因：\n{summary.failure_reason_output_path}"])
        elif self.mode.get() == MODE_MIXED:
            lines = [
                "混合表自动分流完成。",
                f"生效短链：{summary.shortlink_rows:,} 条",
                f"生效签名：{summary.signature_rows:,} 条",
                f"完全重复并去重：{summary.duplicate_rows:,} 条",
                f"异常并跳过：{summary.error_rows:,} 条",
            ]
            lines.append("")
            if summary.shortlink_output_path:
                lines.extend(["", f"生效短链结果：\n{summary.shortlink_output_path}"])
            if summary.signature_output_path:
                lines.extend(["", f"生效签名结果：\n{summary.signature_output_path}"])
        else:
            lines = [
                f"处理完成，共生成 {summary.generated_rows:,} 条数据。",
                f"完全重复并去重：{summary.duplicate_rows:,} 条",
                f"异常并跳过：{summary.error_rows:,} 条",
                "",
                f"结果文件：\n{summary.output_path}",
            ]
        if summary.error_path:
            lines.extend(["", f"错误明细：\n{summary.error_path}"])
        messagebox.showinfo(APP_NAME, "\n".join(lines))


def run_cli(arguments: Sequence[str]) -> int:
    parser = argparse.ArgumentParser(description=APP_NAME)
    parser.add_argument(
        "--mode",
        choices=(
            MODE_SHORTLINK,
            MODE_SIGNATURE,
            MODE_MIXED,
            MODE_SELF_MOBILE,
            MODE_SELF_TELECOM,
            MODE_TABLE_TOOLS,
            MODE_SIGNATURE_CONSISTENCY,
            MODE_PORT_ALLOCATION,
        ),
        required=True,
    )
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--master", type=Path)
    parser.add_argument("--channel", default="")
    parser.add_argument("--start-sub-port", type=int, default=0)
    parser.add_argument("--operator", choices=("移动", "联通", "电信"), required=True)
    parser.add_argument("--business-type", choices=("行业", "营销"), default="行业")
    parser.add_argument("--important-support", choices=("是", "否"), required=True)
    parser.add_argument("--template-dir", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--signature-total-dir", type=Path)
    # 原五种模式允许启动器不传或传入空值；只有拆分/合并模式才在
    # validate_parameters() 中要求 split/merge，避免新增参数拦截旧功能。
    parser.add_argument("--table-action", default="")
    parser.add_argument("--max-rows", type=int, default=1000)
    parser.add_argument(
        "--signature-match-mode",
        choices=(SIGNATURE_MATCH_FULL_RULE, SIGNATURE_MATCH_REGION_ONLY),
        default=SIGNATURE_MATCH_FULL_RULE,
    )
    parser.add_argument("--cancel-file", type=Path)
    args = parser.parse_args(arguments)
    cancel_file = args.cancel_file.resolve() if args.cancel_file else None

    def report_table_progress(phase: str, completed: int, total: int, message: str) -> None:
        safe_message = message.replace("\t", " ").replace("\r", " ").replace("\n", " ")
        print(
            f"@@TABLE_PROGRESS\t{phase}\t{completed}\t{total}\t{safe_message}",
            file=sys.stderr,
            flush=True,
        )

    try:
        summary = process_file(
            args.source,
            args.mode,
            args.channel,
            args.start_sub_port,
            args.operator,
            args.business_type,
            args.important_support,
            args.template_dir,
            status_callback=lambda message: print(message, file=sys.stderr),
            table_action=args.table_action,
            max_rows=args.max_rows,
            table_progress_callback=(
                report_table_progress if args.mode == MODE_TABLE_TOOLS else None
            ),
            table_cancel_check=(
                (lambda: bool(cancel_file and cancel_file.exists()))
                if args.mode == MODE_TABLE_TOOLS
                else None
            ),
            signature_match_mode=args.signature_match_mode,
            master_path=args.master,
            signature_total_dir=args.signature_total_dir,
        )
        summary = relocate_summary_outputs(summary, args.output_dir, args.mode)
    except ProcessingError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    print(json.dumps(asdict(summary), ensure_ascii=False))
    return 0


def main() -> int:
    if "--manage-mapping" in sys.argv:
        args = [argument for argument in sys.argv[1:] if argument not in ("--manage-mapping", "--cli")]
        return run_mapping_cli(args)
    if "--ui-smoke-test" in sys.argv:
        app = ImporterApp()
        app.withdraw()
        app.update_idletasks()
        if app.title() != APP_NAME:
            app.destroy()
            return 1
        app.destroy()
        return 0
    if "--cli" in sys.argv:
        args = [argument for argument in sys.argv[1:] if argument != "--cli"]
        return run_cli(args)
    app = ImporterApp()
    app.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
