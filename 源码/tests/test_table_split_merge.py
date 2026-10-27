import csv
import importlib.util
import re
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


SOURCE_PATH = (
    Path(__file__).resolve().parents[1]
    / "通道生效签名短链工具"
    / "table_split_merge.py"
)
SPEC = importlib.util.spec_from_file_location("table_split_merge_under_test", SOURCE_PATH)
assert SPEC is not None and SPEC.loader is not None
table_tools = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = table_tools
SPEC.loader.exec_module(table_tools)


def save_xlsx(path: Path, rows: list[list[object]], title: str = "0") -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    for row in rows:
        worksheet.append(row)
    worksheet.row_dimensions[1].height = 32
    worksheet.row_dimensions[2].font = Font(name="宋体")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1769AA")
    worksheet.column_dimensions["A"].width = 18
    workbook.save(path)
    workbook.close()


def read_xlsx_rows(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def replace_declared_dimension(path: Path, declared_range: str) -> None:
    temporary = path.with_suffix(".dimension-test.xlsx")
    with ZipFile(path, "r") as source, ZipFile(
        temporary, "w", compression=ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content, replacements = re.subn(
                    br'<dimension\s+ref="[^"]+"\s*/>',
                    f'<dimension ref="{declared_range}"/>'.encode("ascii"),
                    content,
                    count=1,
                )
                if replacements != 1:
                    raise AssertionError("测试工作簿中没有找到 dimension 声明")
            target.writestr(item, content)
    temporary.replace(path)


class TableSplitMergeTest(unittest.TestCase):
    def test_xlsx_split_and_merge_round_trip_preserves_rows_and_header_style(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source = root / "客户总表.xlsx"
            rows = [["编号", "名称"], *[[index, f"客户{index}"] for index in range(1, 8)]]
            save_xlsx(source, rows, title="报备数据")

            split = table_tools.split_table(source, max_rows=4)

            self.assertEqual(split.generated_rows, 7)
            self.assertEqual(split.output_file_count, 3)
            self.assertEqual(
                [Path(path).name for path in split.output_paths],
                [
                    "001_客户总表_输出结果.xlsx",
                    "002_客户总表_输出结果.xlsx",
                    "003_客户总表_输出结果.xlsx",
                ],
            )
            self.assertEqual([len(read_xlsx_rows(Path(path))) for path in split.output_paths], [4, 4, 2])

            first_workbook = load_workbook(split.output_paths[0])
            first_sheet = first_workbook["报备数据"]
            self.assertTrue(first_sheet["A1"].font.bold)
            self.assertEqual(first_sheet["A1"].fill.fgColor.rgb, "001769AA")
            self.assertEqual(first_sheet.row_dimensions[1].height, 32)
            self.assertEqual(first_sheet.row_dimensions[2].font.name, "宋体")
            self.assertEqual(first_sheet.column_dimensions["A"].width, 18)
            first_workbook.close()

            merged = table_tools.merge_tables(
                Path(split.output_dir), now=datetime(2026, 8, 27, 15, 30)
            )
            self.assertEqual(merged.source_file_count, 3)
            self.assertEqual(merged.generated_rows, 7)
            self.assertEqual(
                Path(merged.output_path).name,
                "合并总表_20260827_1530_输出结果.xlsx",
            )
            self.assertEqual(read_xlsx_rows(Path(merged.output_path)), rows)
            merged_workbook = load_workbook(merged.output_path, data_only=False)
            merged_sheet = merged_workbook["报备数据"]
            self.assertTrue(merged_sheet["A1"].font.bold)
            self.assertEqual(merged_sheet["A1"].fill.fgColor.rgb, "001769AA")
            merged_workbook.close()

    def test_fast_xlsx_merge_ignores_incorrect_declared_dimension(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            save_xlsx(root / "001.xlsx", [["编号", "名称"], [1, "甲"], [2, "乙"]])
            save_xlsx(root / "002.xlsx", [["编号", "名称"], [3, "丙"], [4, "丁"]])
            replace_declared_dimension(root / "002.xlsx", "A1:A1")

            merged = table_tools.merge_tables(root)

            self.assertEqual(merged.generated_rows, 4)
            self.assertEqual(
                read_xlsx_rows(Path(merged.output_path)),
                [["编号", "名称"], [1, "甲"], [2, "乙"], [3, "丙"], [4, "丁"]],
            )

    def test_split_cancel_removes_partial_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source = root / "客户总表.xlsx"
            save_xlsx(
                source,
                [["编号"], *[[index] for index in range(1, 8)]],
            )
            cancel = False

            def progress(_phase: str, completed: int, _total: int, _message: str) -> None:
                nonlocal cancel
                if completed > 0:
                    cancel = True

            with self.assertRaises(table_tools.TableOperationCancelled):
                table_tools.split_table(
                    source,
                    max_rows=3,
                    progress_callback=progress,
                    cancel_check=lambda: cancel,
                )

            output_dir = root / table_tools.OUTPUT_DIR_NAME
            self.assertEqual(list(output_dir.glob("*.xlsx")), [])

    def test_fast_merge_cancel_removes_temporary_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            save_xlsx(root / "001.xlsx", [["编号"], [1], [2]])
            save_xlsx(root / "002.xlsx", [["编号"], [3], [4]])
            cancel = False

            def progress(phase: str, completed: int, _total: int, _message: str) -> None:
                nonlocal cancel
                if phase == "merge" and completed > 0:
                    cancel = True

            with self.assertRaises(table_tools.TableOperationCancelled):
                table_tools.merge_tables(
                    root,
                    progress_callback=progress,
                    cancel_check=lambda: cancel,
                )

            output_dir = root / table_tools.OUTPUT_DIR_NAME
            self.assertEqual(list(output_dir.iterdir()), [])

    def test_xlsx_merge_rejects_strict_header_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            save_xlsx(root / "001.xlsx", [["编号", "名称"], [1, "甲"]])
            save_xlsx(root / "002.xlsx", [["编号", "姓名"], [2, "乙"]])

            with self.assertRaisesRegex(table_tools.TableOperationError, "第 2 列"):
                table_tools.merge_tables(root)

    def test_csv_split_merge_supports_gb18030_and_utf8_bom(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source = root / "客户总表.csv"
            source.write_bytes(
                '编号,名称\r\n1,"甲\n公司"\r\n2,乙\r\n3,丙\r\n'.encode("gb18030")
            )

            split = table_tools.split_table(source, max_rows=3)
            self.assertEqual(split.output_file_count, 2)
            for path in split.output_paths:
                self.assertTrue(Path(path).read_bytes().startswith(b"\xef\xbb\xbf"))

            merged = table_tools.merge_tables(
                Path(split.output_dir), now=datetime(2026, 8, 27, 9, 5)
            )
            with Path(merged.output_path).open("r", encoding="utf-8-sig", newline="") as stream:
                self.assertEqual(
                    list(csv.reader(stream)),
                    [["编号", "名称"], ["1", "甲\n公司"], ["2", "乙"], ["3", "丙"]],
                )

    def test_merge_rejects_mixed_formats_and_ignores_previous_merged_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            save_xlsx(root / "001.xlsx", [["编号"], [1]])
            save_xlsx(root / "002.xlsx", [["编号"], [2]])
            save_xlsx(
                root / "合并总表_20260827_1200_输出结果.xlsx",
                [["编号"], [999]],
            )
            first = table_tools.merge_tables(root, now=datetime(2026, 8, 27, 12, 1))
            self.assertEqual(first.source_file_count, 2)
            self.assertEqual(read_xlsx_rows(Path(first.output_path)), [["编号"], [1], [2]])

            (root / "003.csv").write_text("编号\n3\n", encoding="utf-8-sig")
            with self.assertRaisesRegex(table_tools.TableOperationError, "不能混合"):
                table_tools.merge_tables(root)


if __name__ == "__main__":
    unittest.main()
